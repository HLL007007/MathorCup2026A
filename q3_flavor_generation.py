"""第三问改进版：收缩条件高斯联结模型及受约束生成/迁移。
图标题和说明为中文，化合物、感官指标及风格名称保留英文。

运行：python q3_flavor_generation.py；精简检查：添加 --quick。
仅重画已有结果：python q3_flavor_generation.py --plot-only "运行结果/Q3_..."
图片写入该结果目录的 figures_中文；重画会更新同名图片，不修改工作簿数值。
验证和最终工作簿均保留迁移谱、逐项约束与汇总；迁移专表单列成功谱。
感官画像覆盖50项，按香气、风味、口感、其他香气、缺陷五类分别合并展示。

本版取消条件变分模型及torch依赖。使用同一训练分布比较原始抽样与约束抽样。
化学范围由训练内部折外样本校准，不再用拟合样本自身距离设置过窄门槛。
四种缺陷浓度采用截断联合高斯的条件抽样，保持与其余化合物的条件关联，
不独立裁剪浓度。有限步Gibbs为近似抽样；每个最终样本仍逐项复核。
迁移在目标风格候选与原样本间搜索。物理缺陷超阈值程度不得增加；
陈旧酒花仍为感官预测代理。失败及缺失约束不得标为成功。
分类器最高概率类别仅作诊断，风格支持门槛由训练内部折外概率校准。
所有校准均不使用外部留出集；此次迭代已经参考旧版验证结果，留出集称为开发验证。
化学谱沿用原始表示，四种缺陷列使用表8浓度；汇总项是估计附加值，见输出说明。
"""
from __future__ import annotations

import argparse
import hashlib
import importlib
import json
import math
import platform
import re
import textwrap
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
from joblib import Parallel, delayed
from scipy.cluster.hierarchy import cut_tree, fcluster, linkage
from scipy.spatial.distance import cdist, squareform
from scipy.special import ndtri
from scipy.stats import wasserstein_distance, truncnorm, chi2
from sklearn.covariance import LedoitWolf
from sklearn.cross_decomposition import PLSRegression
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression, Ridge
from sklearn.metrics import accuracy_score, balanced_accuracy_score, f1_score, log_loss
from sklearn.preprocessing import StandardScaler

try:
    import q1_without_joint as q1
except ImportError as exc:
    raise ImportError('请将本文件与已完成的 q1_without_joint.py 放在同一目录，并安装第一问依赖。') from exc

MODELS = ('GaussianCopula', 'Copula_constrained')
FOCUS = ('Tripel', 'Kriek', 'Lambic')
DERIVED = ('kcalperc', 'esters_sum', 'aroma_hops_sum', 'acids_sum', 'sulfur_sum')
DEFECTS = ('DMS', 'X2.3.butanedione', 'X4VG', 'X.E..2.nonenal')
DEFECT_LABELS = ('DMS', 'Diacetyl', '4VG', 'trans-2-nonenal')
NONENAL_SOURCE = 'https://www.sciencedirect.com/science/article/pii/0950329393901583'
PROTOCOL = 'Q3_style_copula_constrained_v2'
EPS = 1e-10
MODEL_ZH = {'GaussianCopula': '原始高斯联结模型', 'Copula_constrained': '受约束高斯联结模型',
            'CVAE_screened': '旧版条件变分模型', 'Original': '迁移前',
            'Real_observed': '真实样本', 'Real': '真实样本'}
STYLE_ZH = {'Tripel': '三料啤酒', 'Kriek': '樱桃兰比克', 'Lambic': '兰比克',
    'Pils/Lager': '皮尔森／拉格', 'Amber': '琥珀啤酒', 'Blond': '金色艾尔',
    'Brett/cofermented': '布雷特／混合发酵', 'Brown': '棕色艾尔', 'Brut': '干型啤酒',
    'Christmas': '圣诞啤酒', 'Dubbel': '双料啤酒', 'Faro': '法罗啤酒',
    'Flanders old brown': '弗兰德老棕', 'Fruitbeer': '水果啤酒', 'Hoppy': '酒花型啤酒',
    'Low/No alcohol': '低醇／无醇', 'Saison': '赛松', 'Scotch': '苏格兰艾尔',
    'Stout/Porter': '世涛／波特', 'Strong ale': '烈性艾尔',
    'West Flanders ale': '西弗兰德艾尔', 'Wheat': '小麦啤酒'}
SENSORY_DOMAINS = {
    'Aroma': ('香气', ()), 'Flavour': ('风味', ()),
    'Palate': ('口感与总体评价', ('acidity', 'bitternes', 'sweetness', 'alcohol', 'aftertaste', 'body', 'co2', 'overall')),
    'Other_aroma': ('其他香气', ('orange', 'coriander', 'clove', 'lactic', 'acetic', 'barnyard')),
    'Defects': ('缺陷相关感官', ('X4vg', 'diacetyl', 'dms', 'metallic', 'stale_hops', 't2n'))}


def model_zh(name):
    if name.endswith('_before'):
        return MODEL_ZH.get(name[:-7], '模型')+'：迁移前'
    return MODEL_ZH.get(name, '模型')


def sensory_group_label(members):
    return ' / '.join(str(members).split(' | '))


def chemical_display_name(name, feature):
    # 与第二问相同的希腊字母乱码修正；无名称时保留原字段，不编造化合物身份。
    label = name.strip() if isinstance(name, str) and name.strip() else str(feature)
    return label.replace('Î±', 'alpha').replace('Î²', 'beta').replace('Î³', 'gamma')


def note(message):
    print(message, flush=True)


def serializable(value):
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(type(value).__name__)


def write_json(value, path):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=serializable), encoding='utf-8')


def safe_name(value):
    return re.sub(r'[\\/:*?"<>|]+', '_', str(value))


def check_dependencies(plots=True):
    for package in ('openpyxl',) + (('matplotlib',) if plots else ()):
        try:
            importlib.import_module(package)
        except ImportError as exc:
            raise ImportError(f'当前Python缺少 {package}，请在同一环境运行：python -m pip install {package}') from exc


def save_book(path, tables):
    """用户运行时用openpyxl导出；原输入不覆盖，临时文件完成后替换本次输出。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.stem + '.writing.xlsx')
    with pd.ExcelWriter(temp, engine='openpyxl') as writer:
        for name, frame in tables.items():
            if len(frame) > 1048575 or len(frame.columns) > 16384:
                raise ValueError(f'工作表超过Excel容量：{name}')
            clean = frame.replace([np.inf, -np.inf], np.nan).copy()
            # 防止文本字段被Excel解释为公式；只处理字符串，不改变负数。
            for column in clean.select_dtypes(include=['object', 'string']).columns:
                clean[column] = clean[column].map(
                    lambda v: "'" + v if isinstance(v, str) and v.startswith(('=', '+', '@')) else v)
            clean.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            for cell in ws[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='264653')
            for j, col in enumerate(clean.columns, 1):
                length = max([len(str(col))] + [len(str(v)) for v in clean[col].head(30)])
                ws.column_dimensions[get_column_letter(j)].width = min(48, max(12, length + 2))
    temp.replace(path)


def stratified_holdout(styles, fraction, seed):
    rng = np.random.default_rng(seed)
    train, valid = [], []
    for style in sorted(set(styles)):
        ids = rng.permutation(np.flatnonzero(styles == style))
        if len(ids) < 3:
            raise ValueError(f'风格 {style} 少于3个样本，无法执行当前留出协议。')
        nv = max(1, min(len(ids)-2, int(round(len(ids)*fraction))))
        valid.extend(ids[:nv])
        train.extend(ids[nv:])
    return np.sort(train), np.sort(valid)


def balanced_folds(styles, count, seed):
    rng = np.random.default_rng(seed)
    buckets, offset = [[] for _ in range(count)], 0
    for style in rng.permutation(np.unique(styles)):
        ids = rng.permutation(np.flatnonzero(styles == style))
        for j, idx in enumerate(ids):
            buckets[(offset+j) % count].append(int(idx))
        offset = (offset+len(ids)) % count
    all_ids = np.arange(len(styles))
    if any(len(b) < 2 for b in buckets):
        raise ValueError('内部折样本不足。')
    return [(np.setdiff1d(all_ids, b), np.sort(b)) for b in buckets]


def choose_q1_family(root, provided, override):
    if override != 'auto':
        return override, None
    folders = [provided] if provided else sorted((root/'运行结果').glob('*'), reverse=True)
    valid = []
    for folder in folders:
        if folder and (folder/'best_model.json').is_file() and (folder/'run_manifest.json').is_file():
            m = json.loads((folder/'run_manifest.json').read_text(encoding='utf-8-sig'))
            b = json.loads((folder/'best_model.json').read_text(encoding='utf-8-sig'))
            if m.get('status') == 'completed' and b.get('model') in q1.MODELS:
                valid.append((folder.stat().st_mtime, b['model'], folder))
    if not valid:
        raise FileNotFoundError('未找到完成的第一问最佳模型记录。可指定 --q1-results 路径，或 --forward-model TargetBlend。')
    _, family, folder = max(valid, key=lambda item: item[0])
    return family, folder


def load_data(args):
    root = q1.resolve_data_root(args.data_root)
    x, y, meta, names, targets, files = q1.load_arrays(root, 'all')
    # pandas的Copy-on-Write可能让to_numpy返回只读视图；校准替换只操作独立副本。
    x = np.array(x, dtype=float, copy=True)
    final = root/'处理结果/最终数据'
    files.update(reference=final/'8_参考数据.xlsx', sensory_correlation=final/'5_感官相关矩阵.csv',
                 field_metadata=root/'处理结果/处理记录/结构整理（第4版）/05_化合物字段说明.csv')
    ref = pd.read_excel(files['reference'], sheet_name='Sensory thresholds')
    ref['field_name_R'] = ref.field_name_R.astype(str).str.strip()
    ref = ref.loc[ref.field_name_R.isin(names)].copy()
    if ref.field_name_R.duplicated().any() or set(ref.field_name_R) != set(names):
        raise ValueError('表8化合物说明不能与231个字段一一对应。')
    ref = ref.set_index('field_name_R').loc[names]
    calibration = pd.read_excel(files['reference'], sheet_name='Calibrated chemical dataset')
    if calibration.beer.duplicated().any() or set(calibration.beer) != set(meta.beer):
        raise ValueError('表8校准数据beer编号不一致。')
    calibration = calibration.set_index('beer').loc[meta.beer]
    if not np.array_equal(calibration.tasting_category_fine.astype(str), meta.tasting_category_fine.astype(str)):
        raise ValueError('表8与主数据风格标签不一致。')
    raw = x.copy()
    changes = []
    for field in DEFECTS:
        j = names.index(field)
        values = pd.to_numeric(calibration[field], errors='raise').to_numpy(float)
        if np.isinf(values).any() or np.any(values[np.isfinite(values)] < 0):
            raise ValueError(f'校准浓度必须非负且非无穷：{field}')
        for i in range(len(x)):
            changes.append(dict(beer=int(meta.beer.iloc[i]), feature=field, original_value=raw[i, j],
                                calibrated_concentration=values[i], unit=ref.loc[field, 'unit'],
                                missing_calibration=not np.isfinite(values[i])))
        x[:, j] = values  # 不回写输入文件；缺失继续保留，由训练内填补。
    internal = set(ref.index[ref.origin_paper.astype(str).str.contains('internal standard', case=False)])
    excluded = set(DERIVED) | internal
    keep = np.array([n not in excluded for n in names])
    chemical_names = list(np.array(names)[keep])
    bounds = ref.loc[list(DEFECTS), 'threshold_median'].to_numpy(dtype=float, copy=True)
    bounds[-1] = args.nonenal_threshold_mg_l
    if not np.all(np.isfinite(bounds) & (bounds > 0)):
        raise ValueError('缺陷阈值缺失或非正。')
    if list(ref.loc[list(DEFECTS), 'unit']) != ['ug/L', 'mg/L', 'mg/L', 'mg/L']:
        raise ValueError('缺陷浓度单位与约束定义不符，拒绝自动换算未知单位。')
    field_meta = pd.read_csv(files['field_metadata']).set_index('field_name_R').loc[names]
    catalog = pd.DataFrame({'feature': names, 'chemical_name': ref.chemical_name.to_numpy(),
        'unit': ref.unit.to_numpy(), 'original_transform': field_meta['候选关系'].to_numpy(),
        'original_verification': field_meta['核验状态'].to_numpy(),
        'model_input': keep,
        'output_representation': ['table8 calibrated physical concentration' if n in DEFECTS else
            'estimated aggregate in supplied representation' if n in DERIVED else
            'fixed training median; internal standard' if n in internal else
            'table1 supplied representation; may be logarithmic' for n in names]})
    catalog['图中名称'] = [chemical_display_name(r.chemical_name, r.feature) for r in catalog.itertuples()]
    r = pd.read_csv(files['sensory_correlation'], index_col=0)
    if set(r.index) != set(targets) or set(r.columns) != set(targets):
        raise ValueError('表5感官字段与表4不一致。')
    r = r.loc[targets, targets].to_numpy(float)
    if not (np.isfinite(r).all() and np.max(np.abs(r)) <= 1+1e-8 and np.allclose(r, r.T) and np.allclose(np.diag(r), 1)):
        raise ValueError('表5不是有效的对称相关矩阵。')
    groups = radar_groups(r, targets, args.radar_min_corr)
    if meta.tasting_category_fine.isna().any():
        raise ValueError('风格标签缺失。')
    styles = meta.tasting_category_fine.astype(str).to_numpy()
    requested = sorted(set(styles)) if args.styles is None else list(dict.fromkeys(args.styles))
    if set(requested) - set(styles):
        raise ValueError(f'未知风格：{set(requested)-set(styles)}')
    # 三个重点风格始终交付；迁移目标Tripel同样在生成列表。
    requested = list(dict.fromkeys(requested + list(FOCUS)))
    family, q1_folder = choose_q1_family(root, args.q1_results, args.forward_model)
    return SimpleNamespace(root=root, x=x[:, keep], x_full=x, raw=raw, y=y, meta=meta,
        names=chemical_names, all_names=names, targets=targets, styles=styles, files=files,
        catalog=catalog, calibration_changes=pd.DataFrame(changes), thresholds=bounds,
        threshold_table=pd.DataFrame({'feature': DEFECTS, 'defect': DEFECT_LABELS,
            'threshold': bounds, 'unit': ['ug/L', 'mg/L', 'mg/L', 'mg/L'],
            'source': ['table8 threshold_median']*3+[NONENAL_SOURCE],
            'scope': ['sensory threshold; not food-safety limit']*4}),
        groups=groups, requested=requested, internal=internal, family=family, q1_folder=q1_folder)


def radar_groups(r, targets, minimum):
    rows = []
    for domain, (_, fields) in SENSORY_DOMAINS.items():
        if domain in ('Aroma', 'Flavour'):
            prefix = 'A_' if domain == 'Aroma' else 'F_'
            ids = [j for j, n in enumerate(targets) if n.startswith(prefix)]
        else:
            ids = [j for j, n in enumerate(targets) if n in fields]
        if not ids:
            continue
        distance = np.clip(1-r[np.ix_(ids, ids)], 0, 2)
        np.fill_diagonal(distance, 0)
        labels = np.arange(len(ids))
        if len(ids) > 1:
            tree = linkage(squareform(distance), method='complete')
            labels = fcluster(tree, 1-minimum, criterion='distance')
            if len(set(labels)) < min(3, len(ids)):
                # 仅拆分过大的相关组，不降低相关阈值；雷达至少保留三个可解释轴。
                labels = cut_tree(tree, n_clusters=[min(3, len(ids))]).ravel()
        ordered = sorted(set(labels), key=lambda k: min(j for j, v in enumerate(labels) if v == k))
        for number, group in enumerate(ordered, 1):
            member_ids = [ids[j] for j, v in enumerate(labels) if v == group]
            members = [targets[j] for j in member_ids]
            correlations = r[np.ix_(member_ids, member_ids)][np.triu_indices(len(member_ids), 1)]
            actual_minimum = float(correlations.min()) if len(correlations) else np.nan
            if len(correlations) and actual_minimum < minimum-1e-8:
                raise ValueError('雷达分组包含低于指定正相关阈值的指标。')
            label = sensory_group_label(' | '.join(members))
            for member in members:
                rows.append(dict(domain=domain, axis=number, label=label, target=member,
                                 members=' | '.join(members), min_pairwise_rho=actual_minimum,
                                 correlation_threshold=minimum,
                                 rule='complete-linkage within domain; positive correlation; at least 3 axes when available'))
    result = pd.DataFrame(rows)
    if result.empty or result.target.duplicated().any() or set(result.target) != set(targets):
        raise ValueError('雷达分组必须恰好覆盖全部感官指标一次。')
    return result


class NormalScores:
    """训练经验边际的单调正态分数变换，不降低维度；逆变换不外推观测范围。"""
    def fit(self, x):
        self.imputer = SimpleImputer(strategy='median', keep_empty_features=True).fit(x)
        a = self.imputer.transform(x)
        self.curves = []
        for col in a.T:
            unique, counts = np.unique(col, return_counts=True)
            probability = (np.cumsum(counts)-counts/2)/len(col)
            self.curves.append((unique, ndtri(np.clip(probability, 1e-5, 1-1e-5))))
        return self

    def transform(self, x):
        a = self.imputer.transform(np.asarray(x, float))
        return np.column_stack([np.interp(a[:, j], u, v) if len(u)>1 else np.zeros(len(a))
                                for j, (u, v) in enumerate(self.curves)])

    def inverse(self, z):
        return np.column_stack([np.interp(z[:, j], v, u) if len(u)>1 else np.full(len(z), u[0])
                                for j, (u, v) in enumerate(self.curves)])

    def upper_bounds(self, columns, concentrations):
        indices, bounds = [], []
        for j, concentration in zip(columns, concentrations):
            u, v = self.curves[j]
            if not np.isfinite(concentration):
                raise SamplingInfeasible('缺陷浓度上限未知，不能认证约束。')
            if concentration < u[0]-1e-12:
                raise SamplingInfeasible('缺陷上限低于训练数据支持范围，当前模型不能生成可行谱。')
            if concentration >= u[-1]:
                continue
            indices.append(j)
            bounds.append(np.nextafter(np.interp(concentration, u, v), -np.inf))
        return np.array(indices, dtype=int), np.array(bounds, dtype=float)


class SamplingInfeasible(ValueError):
    pass


def constrained_gaussian(mean, chol, n, rng, indices=None, bounds=None, sweeps=60):
    """先近似抽取截断的缺陷边际，再条件抽取其余维度；不逐列裁剪化学值。"""
    samples = mean+rng.normal(size=(n, len(mean)))@chol.T
    if indices is None or len(indices) == 0:
        return samples
    covariance = chol@chol.T
    block = covariance[np.ix_(indices, indices)]
    precision = np.linalg.inv(block)
    mu = mean[indices]
    deviation = np.sqrt(np.diag(block))
    # 多条独立链，使用同样的固定迭代预算；近似性在输出方法说明中声明。
    constrained = np.tile(np.minimum(mu, bounds-.1*deviation), (n, 1))
    for _ in range(sweeps):
        for j in rng.permutation(len(indices)):
            delta = constrained-mu
            conditional_mean = mu[j]-(delta@precision[j]-delta[:, j]*precision[j, j])/precision[j, j]
            conditional_sd = np.sqrt(1/precision[j, j])
            standardized_upper = (bounds[j]-conditional_mean)/conditional_sd
            constrained[:, j] = truncnorm.rvs(-np.inf, standardized_upper, loc=conditional_mean,
                                               scale=conditional_sd, size=n, random_state=rng)
    gain = np.linalg.solve(block, covariance[:, indices].T).T
    samples += (constrained-samples[:, indices])@gain.T
    samples[:, indices] = constrained  # 已条件抽到的值，消除浮点求解残差，不是裁剪。
    if not np.isfinite(samples).all() or np.any(samples[:, indices] > bounds+1e-9):
        raise FloatingPointError('截断联合抽样出现非有限值或越界。')
    return samples


class ConditionalCopula:
    def fit(self, x, styles, mean_strength=5., covariance_strength=20.):
        self.map = NormalScores().fit(x)
        z = self.map.transform(x)
        self.global_mean = z.mean(0)
        residual = z.copy()
        self.parameters = {}
        for style in np.unique(styles):
            residual[styles == style] -= z[styles == style].mean(0)
        shared = LedoitWolf().fit(residual).covariance_
        for style in np.unique(styles):
            a = z[styles == style]
            n = len(a)
            mean = (n*a.mean(0)+mean_strength*self.global_mean)/(n+mean_strength)
            covariance = np.cov(a, rowvar=False) if n>1 else shared
            weight = n/(n+covariance_strength)
            covariance = weight*covariance+(1-weight)*shared+np.eye(z.shape[1])*1e-5
            self.parameters[style] = (mean, np.linalg.cholesky(covariance))
        self.train_z = z
        return self

    def sample(self, style, n, rng, source=None, source_style=None, defect_ids=None, caps=None, sweeps=60):
        mean, chol = self.parameters[style]
        indices, upper = self.map.upper_bounds(defect_ids, caps) if caps is not None else (None, None)
        z = constrained_gaussian(mean, chol, n, rng, indices, upper, sweeps)
        if source is not None:
            # 在正态分数空间搜索迁移路径；源谱与受约束终点满足同一上界时，
            # 其凸组合也满足上界。候选仍需通过全部正向检查，不保证找到可行解。
            source_z = self.map.transform(source[None])[0]
            strength = rng.choice([.35, .55, .75, 1.], size=n)[:, None]
            z = source_z+strength*(z-source_z)
        return self.map.inverse(z)

    def distance(self, x, style):
        mean, chol = self.parameters[style]
        residual = np.linalg.solve(chol, (self.map.transform(x)-mean).T).T
        return np.sqrt(np.mean(residual**2, axis=1))




def fit_base(name, parameters, x, y, jobs, seed):
    if name == 'Ridge':
        alpha = np.array([p['alpha'] for p in parameters]) if isinstance(parameters, list) else parameters['alpha']
        return Ridge(alpha=alpha, solver='svd').fit(x, y)
    if name == 'PLS2':
        k = min(parameters['n_components'], len(x)-1, x.shape[1])
        return PLSRegression(n_components=max(1, k), scale=False, max_iter=1500).fit(x, y)
    def one(j):
        pars = parameters[j] if isinstance(parameters, list) else parameters
        model = RandomForestRegressor(**pars, n_jobs=1, random_state=seed+j) if name == 'RF' else \
            GradientBoostingRegressor(**pars, random_state=seed+j)
        return model.fit(x, y[:, j])
    return Parallel(n_jobs=jobs, prefer='threads')(delayed(one)(j) for j in range(y.shape[1]))


def base_predict(model, x, jobs):
    if isinstance(model, list):
        return np.column_stack(Parallel(n_jobs=jobs, prefer='threads')(delayed(m.predict)(x) for m in model))
    return np.asarray(model.predict(x))


class ForwardModel:
    def fit(self, x, y, styles, targets, family, args, seed):
        self.family, self.jobs = family, args.jobs
        bases = list(q1.BASE_MODELS) if family in q1.FUSION_MODELS else [family]
        folds = balanced_folds(styles, args.forward_folds, seed)
        selected, oofs, records = {}, [], []
        for name in bases:
            note(f'  感官模型内部选参：{name}')
            candidates = q1.make_candidates(name, args.quick)
            predictions = np.zeros((len(candidates), *y.shape))
            for h, (tr, va) in enumerate(folds):
                pre = q1.FoldPreprocessor(.01).fit(x[tr])
                sy = StandardScaler().fit(y[tr])
                xt, xv, yt = pre.transform(x[tr]), pre.transform(x[va]), sy.transform(y[tr])
                for k, parameters in enumerate(candidates):
                    fitted = fit_base(name, parameters, xt, yt, args.jobs, seed+h)
                    predictions[k, va] = sy.inverse_transform(base_predict(fitted, xv, args.jobs))
            loss = np.mean(np.abs(predictions-y[None]), axis=1)
            if name == 'PLS2':
                index = int(np.argmin(loss.mean(1)))
                parameters = candidates[index]
                chosen = predictions[index]
            else:
                index = np.argmin(loss, axis=0)
                parameters = [candidates[i] for i in index]
                chosen = np.column_stack([predictions[index[j], :, j] for j in range(y.shape[1])])
            selected[name] = parameters
            oofs.append(chosen)
            records.append(dict(model=name, parameters=json.dumps(parameters), inner_MAE=float(np.mean(np.abs(chosen-y)))))
        if family in q1.FUSION_MODELS:
            learned, weight_records = q1.learn_target_fusion(y, np.stack(oofs, axis=2), targets, .05)
            self.weights = learned[family]
            self.weight_records = pd.DataFrame(weight_records)
            self.oof = np.sum(np.stack(oofs, axis=2)*self.weights[None], axis=2)
        else:
            self.weights = np.ones((y.shape[1], 1))
            self.weight_records = pd.DataFrame({'target': targets, 'selected_model': family})
            self.oof = oofs[0]
        self.pre = q1.FoldPreprocessor(.01).fit(x)
        self.sy = StandardScaler().fit(y)
        xt, yt = self.pre.transform(x), self.sy.transform(y)
        self.models = [fit_base(n, selected[n], xt, yt, args.jobs, seed+100) for n in bases]
        self.parameters = pd.DataFrame(records)
        self.cuts = q1.fit_grade_cuts(y)
        # 审计器不参与筛选；模型家族固定，训练样本相同，因此仍非外部实验验证。
        self.auditor = Ridge(alpha=100., solver='svd').fit(xt, yt)
        return self

    def predict(self, x):
        z = self.pre.transform(x)
        p = np.stack([self.sy.inverse_transform(base_predict(m, z, self.jobs)) for m in self.models], axis=2)
        result = np.sum(p*self.weights[None], axis=2)
        if not np.isfinite(result).all():
            raise ValueError('感官预测非有限。')
        return result

    def audit(self, x):
        return self.sy.inverse_transform(self.auditor.predict(self.pre.transform(x)))


def sensory_metrics(y, pred, targets, phase, model):
    _, per = q1.metric_values(y, pred)
    return pd.DataFrame({'phase': phase, 'model': model, 'target': targets, **per})


def reference_profile(y, styles, targets):
    rows = []
    for style in np.unique(styles):
        a = y[styles == style]
        for j, target in enumerate(targets):
            rows.append(dict(style=style, target=target, n=len(a), mean=a[:, j].mean(),
                median=np.median(a[:, j]), p10=np.quantile(a[:, j], .1), p90=np.quantile(a[:, j], .9)))
    return pd.DataFrame(rows)


def upper_empirical_quantile(values, coverage):
    a = np.sort(np.asarray(values, float))
    a = a[np.isfinite(a)]
    if not len(a):
        raise ValueError('缺少有效的范围校准样本。')
    index = min(len(a)-1, max(0, math.ceil((len(a)+1)*coverage)-1))
    return float(a[index])


def calibrate_support(x, styles, args, seed):
    """训练内部交叉拟合距离/近邻/风格概率；不声称具有严格共形覆盖保证。"""
    rows = []
    for fold, (train, valid) in enumerate(balanced_folds(styles, 3, seed), 1):
        fitted = ConditionalCopula().fit(x[train], styles[train], args.mean_shrinkage, args.cov_shrinkage)
        classifier = LogisticRegression(C=1., max_iter=3000).fit(fitted.train_z, styles[train])
        z = fitted.map.transform(x[valid])
        probabilities = classifier.predict_proba(z)
        neighbors = cdist(z, fitted.train_z).min(1)/np.sqrt(x.shape[1])
        for style in np.unique(styles[valid]):
            positions = np.flatnonzero(styles[valid] == style)
            distances = fitted.distance(x[valid[positions]], style)
            for pos, distance in zip(positions, distances):
                probability = probabilities[pos, list(classifier.classes_).index(style)]
                rows.append(dict(row_index=int(valid[pos]), fold=fold, style=style,
                    chemical_distance=float(distance), nearest_train_distance=float(neighbors[pos]),
                    negative_log_style_probability=float(-np.log(max(probability, 1e-12)))))
    table = pd.DataFrame(rows)
    theoretical = float(np.sqrt(chi2.ppf(args.support_coverage, x.shape[1])/x.shape[1]))
    chemical_limit = max(theoretical, upper_empirical_quantile(table.chemical_distance, args.support_coverage))
    neighbor_limit = upper_empirical_quantile(table.nearest_train_distance, args.support_coverage)
    style_limit = upper_empirical_quantile(table.negative_log_style_probability, args.support_coverage)
    limits = pd.DataFrame([
        {'指标': '化合物联合距离', '门槛': chemical_limit, '理论高斯半径': theoretical},
        {'指标': '最近真实训练样本距离', '门槛': neighbor_limit},
        {'指标': '目标风格负对数概率', '门槛': style_limit}])
    limits['校准范围'] = '仅训练内部折外样本；全风格合并校准'
    limits['经验覆盖设置'] = args.support_coverage
    return table, limits, chemical_limit, neighbor_limit, style_limit


class Context:
    def __init__(self, data, ids, args, phase, seed):
        self.data, self.args, self.phase, self.seed = data, args, phase, seed
        self.ids = ids
        self.x, self.y, self.styles = data.x[ids], data.y[ids], data.styles[ids]
        self.forward = ForwardModel().fit(self.x, self.y, self.styles, data.targets, data.family, args, seed)
        self.copula = ConditionalCopula().fit(self.x, self.styles, args.mean_shrinkage, args.cov_shrinkage)
        note('  用训练内部折外样本校准化学范围与风格支持门槛')
        (self.calibration, self.calibration_limits, chemical_limit, self.neighbor_limit,
         self.style_limit) = calibrate_support(self.x, self.styles, args, seed+200)
        self.calibration['beer'] = data.meta.beer.iloc[ids[self.calibration.row_index.to_numpy(int)]].to_numpy()
        self.sampling_diagnostics = []
        self.classifier = LogisticRegression(C=1., max_iter=3000).fit(self.copula.train_z, self.styles)
        self.y_scale = np.std(self.y, axis=0)
        self.y_scale[self.y_scale < 1e-8] = 1.
        self.af = np.array([j for j, t in enumerate(data.targets) if t.startswith(('A_', 'F_'))])
        self.stale = data.targets.index('stale_hops')
        self.defect_ids = [data.names.index(n) for n in DEFECTS]
        self.ref = reference_profile(self.y, self.styles, data.targets)
        self.centers, self.ranges, self.flavor_limits, self.chemical_limits = {}, {}, {}, {}
        for style in np.unique(self.styles):
            mask = self.styles == style
            a = self.y[mask]
            center = np.mean(a, axis=0)
            self.centers[style] = center
            lo, hi = np.quantile(a, [.1, .9], axis=0)
            self.ranges[style] = (lo, hi)
            d = np.mean(np.abs((self.forward.oof[mask][:, self.af]-center[self.af])/self.y_scale[self.af]), axis=1)
            # 内部OOF只用于设定筛选规则；外部留出集用于评价。
            self.flavor_limits[style] = max(.1, np.quantile(d, .95)*args.flavor_tolerance)
            self.chemical_limits[style] = chemical_limit
        self._train_full = data.x_full[ids]
        self._full_fill = SimpleImputer(strategy='median', keep_empty_features=True).fit_transform(self._train_full)

    def defect_caps(self, style, source=None):
        if source is not None:
            return np.maximum(source[self.defect_ids], self.data.thresholds)
        reference = self.x[self.styles == style][:, self.defect_ids]
        if self.args.defect_policy == 'below-threshold':
            return self.data.thresholds.copy()
        return np.maximum(self.data.thresholds, np.nanquantile(reference, self.args.defect_quantile, axis=0))

    def profile_distance(self, pred, style):
        return np.mean(np.abs((pred[:, self.af]-self.centers[style][self.af])/self.y_scale[self.af]), axis=1)

    def spectrum(self, x):
        """补齐231列的展示表；汇总列标为估计值，不送入预测或约束。"""
        full = np.empty((len(x), len(self.data.all_names)))
        nearest = cdist(self.copula.map.transform(x), self.copula.train_z)
        k = min(5, len(self.x))
        ids = np.argsort(nearest, axis=1)[:, :k]
        weights = 1/(np.take_along_axis(nearest, ids, axis=1)+1e-6)
        weights /= weights.sum(1, keepdims=True)
        for j, name in enumerate(self.data.all_names):
            if name in self.data.names:
                full[:, j] = x[:, self.data.names.index(name)]
            elif name in DERIVED:
                full[:, j] = np.sum(self._full_fill[ids, j]*weights, axis=1)
            else:
                full[:, j] = np.median(self._full_fill[:, j])
        return full

    def assess(self, x, style, source=None):
        pred = self.forward.predict(x)
        audit = self.forward.audit(x)
        z = self.copula.map.transform(x)
        all_probability = self.classifier.predict_proba(z)
        probability = all_probability[:, list(self.classifier.classes_).index(style)]
        predicted_style = self.classifier.classes_[np.argmax(all_probability, axis=1)]
        style_match = predicted_style == style
        negative_log_probability = -np.log(np.maximum(probability, 1e-12))
        style_supported = negative_log_probability <= self.style_limit
        closest = cdist(z, self.copula.train_z).min(1)/np.sqrt(x.shape[1])
        chemical_distance = self.copula.distance(x, style)
        chemical_ok = np.isfinite(x).all(1) & (closest <= self.neighbor_limit) & (chemical_distance <= self.chemical_limits[style])
        concentration = x[:, self.defect_ids]
        if source is None:
            caps = self.defect_caps(style)
            stale_cap = np.quantile(self.y[self.styles == style, self.stale], self.args.defect_quantile)
            source_known = True
        else:
            caps = self.defect_caps(style, source)
            source_known = bool(np.isfinite(caps).all())
            stale_cap = float(self.forward.predict(source[None])[0, self.stale])
        # 未知约束按失败处理，绝不把未知解释为通过。
        each_pass = (concentration <= caps[None]+1e-9) & np.isfinite(caps[None])
        chemical_defects_ok = each_pass.all(1) & source_known
        stale_ok = pred[:, self.stale] <= stale_cap+1e-8
        flavor = self.profile_distance(pred, style)
        flavor_ok = flavor <= self.flavor_limits[style]
        low, high = self.ranges[style]
        interval = ((pred[:, self.af] >= low[self.af]) & (pred[:, self.af] <= high[self.af])).mean(1)
        score = flavor/self.flavor_limits[style]+.25*(1-probability)+.1*chemical_distance
        changes = np.zeros(len(x))
        source_flavor, source_probability, source_chemical = np.nan, np.nan, np.nan
        improved = np.ones(len(x), bool)
        if source is not None:
            source_z = self.copula.map.transform(source[None])
            changes = np.mean(np.abs(z-source_z), axis=1)
            source_flavor = float(self.profile_distance(self.forward.predict(source[None]), style)[0])
            source_probability = float(self.classifier.predict_proba(source_z)[0, list(self.classifier.classes_).index(style)])
            source_chemical = float(self.copula.distance(source[None], style)[0])
            improved = (flavor < source_flavor-EPS) & (chemical_distance < source_chemical-EPS)
            score += self.args.migration_cost*changes
        table = pd.DataFrame({'flavor_distance_train_reference': flavor, 'target_style_probability': probability,
            'predicted_style': predicted_style, 'style_match': style_match,
            'style_supported': style_supported, 'negative_log_style_probability': negative_log_probability,
            'style_score_limit': self.style_limit,
            'target_interval_fraction': interval, 'chemical_distance': chemical_distance,
            'nearest_train_distance': closest, 'chemical_distance_limit': self.chemical_limits[style],
            'nearest_distance_limit': self.neighbor_limit, 'chemical_ok': chemical_ok, 'flavor_ok': flavor_ok,
            'chemical_defects_ok': chemical_defects_ok, 'stale_proxy_ok': stale_ok,
            'stale_proxy': pred[:, self.stale], 'stale_proxy_cap': stale_cap,
            'defect_source_known': source_known, 'modification_normal_score': changes,
            'source_flavor_distance': source_flavor, 'source_target_probability': source_probability,
            'source_chemical_distance': source_chemical, 'migration_improved': improved,
            'screen_score': score,
            'all_known_concentrations_below_threshold': (concentration <= self.data.thresholds).all(1),
            'audit_flavor_distance': self.profile_distance(audit, style)})
        table['constraint_pass'] = chemical_ok & chemical_defects_ok & stale_ok
        table['eligible'] = table.constraint_pass & flavor_ok & improved & style_supported
        for j, field in enumerate(DEFECTS):
            table[field+'_concentration'] = concentration[:, j]
            table[field+'_threshold'] = self.data.thresholds[j]
            table[field+'_cap'] = caps[j]
            table[field+'_excess'] = np.maximum(concentration[:, j]/self.data.thresholds[j]-1, 0)
            table[field+'_pass'] = each_pass[:, j]
        return table, pred


def diversify(x, eligible_ids, score, count, mapper):
    if not len(eligible_ids):
        return np.array([], dtype=int)
    ids = eligible_ids[np.argsort(score[eligible_ids], kind='stable')]
    ids = ids[:min(len(ids), max(count*5, count))]
    z = mapper.transform(x[ids])
    chosen, remaining = [0], np.ones(len(ids), bool)
    remaining[0] = False
    nearest = np.linalg.norm(z-z[0], axis=1)/np.sqrt(z.shape[1])
    ranked = score[ids]
    normalized_score = (ranked-ranked.min())/(np.ptp(ranked)+EPS)
    while len(chosen) < min(count, len(ids)):
        remaining &= nearest >= 1e-7
        if not remaining.any():
            break  # 仅在剩余候选均重复时停止，不让高排名重复谱挡住其他可行谱。
        utility = nearest-.25*normalized_score
        utility[~remaining] = -np.inf
        j = int(np.argmax(utility))
        chosen.append(j)
        remaining[j] = False
        nearest = np.minimum(nearest, np.linalg.norm(z-z[j], axis=1)/np.sqrt(z.shape[1]))
    return ids[chosen]


def pack(ctx, x, pred, assess, model, style, source_beer=None, status='generated'):
    n = len(x)
    if pred.shape != (n, len(ctx.data.targets)) or len(assess) != n:
        raise ValueError('化合物谱、感官预测和约束检查的行数或目标数不一致。')
    if not np.isfinite(pred).all() or (model != 'Original' and not np.isfinite(x).all()):
        raise ValueError('生成谱或感官预测含非有限值，拒绝导出。')
    if (status == 'successful_migration' or (model == MODELS[1] and status == 'generated')) and not assess.eligible.all():
        raise ValueError('受约束交付或成功迁移中混入了不合格样本。')
    prefix = f'{model}_{safe_name(style)}' if source_beer is None else f'{model}_beer{source_beer}'
    meta = pd.DataFrame({'sample_id': [f'{prefix}_{i+1:04d}' for i in range(n)], 'model': [model]*n,
                         'style': [style]*n, 'source_beer': [source_beer]*n, 'status': [status]*n})
    grade = q1.quantile_grades(pred, ctx.forward.cuts) if n else np.empty((0, len(ctx.data.targets)))
    full = ctx.spectrum(x) if n else np.empty((0, len(ctx.data.all_names)))
    if model == 'Original':
        source_index = np.flatnonzero(ctx.data.meta.beer.to_numpy() == source_beer)
        if n != 1 or len(source_index) != 1:
            raise ValueError('原始迁移样本编号不唯一。')
        # 原样本保留真实汇总项及内标值，不能用生成谱的附加字段估计替换。
        full = ctx.data.x_full[source_index].copy()
    return dict(spectra=pd.concat([meta, pd.DataFrame(full, columns=ctx.data.all_names)], axis=1),
        continuous=pd.concat([meta, pd.DataFrame(pred, columns=ctx.data.targets)], axis=1),
        grades=pd.concat([meta, pd.DataFrame(grade, columns=ctx.data.targets)], axis=1),
        checks=pd.concat([meta, assess.reset_index(drop=True)], axis=1))


def combine(items):
    return {k: pd.concat([v[k] for v in items], ignore_index=True) for k in ('spectra', 'continuous', 'grades', 'checks')}


def evaluate_pool(ctx, x, pred, checks, style, model, stage, reference_x, reference_y, requested):
    row = dict(phase=ctx.phase, style=style, model=model, stage=stage, n_generated=len(x), requested=requested,
               real_reference_n=len(reference_y), reference_role='heldout_real' if ctx.phase=='validation' else 'training_descriptive')
    if not len(x):
        return row
    row.update(flavor_distance=float(checks.flavor_distance_train_reference.mean()),
        target_interval_fraction=float(checks.target_interval_fraction.mean()),
        target_style_probability=float(checks.target_style_probability.mean()),
        target_style_match_rate=float(checks.style_match.mean()),
        style_support_pass_rate=float(checks.style_supported.mean()),
        chemical_support_pass_rate=float(checks.chemical_ok.mean()),
        flavor_pass_rate=float(checks.flavor_ok.mean()),
        eligible_rate=float(checks.eligible.mean()),
        constraint_pass_rate=float(checks.constraint_pass.mean()),
        chemical_defect_pass_rate=float(checks.chemical_defects_ok.mean()),
        stale_proxy_pass_rate=float(checks.stale_proxy_ok.mean()),
        all_below_threshold_rate=float(checks.all_known_concentrations_below_threshold.mean()),
        audit_flavor_distance=float(checks.audit_flavor_distance.mean()))
    z = ctx.copula.map.transform(x)
    if len(x)>1:
        distance = cdist(z, z)/np.sqrt(z.shape[1])
        row['diversity_mean_distance'] = float(distance[np.triu_indices(len(x), 1)].mean())
    row['nearest_train_distance_mean'] = float(checks.nearest_train_distance.mean())
    if len(reference_y):
        row['profile_MAE_standardized'] = float(np.mean(np.abs((pred[:, ctx.af].mean(0)-reference_y[:, ctx.af].mean(0))/ctx.y_scale[ctx.af])))
        row['sensory_Wasserstein_standardized'] = float(np.mean([
            wasserstein_distance(pred[:, j]/ctx.y_scale[j], reference_y[:, j]/ctx.y_scale[j]) for j in ctx.af]))
        rz = ctx.copula.map.transform(reference_x)
        row['chemical_Wasserstein_normal_score'] = float(np.mean([wasserstein_distance(z[:, j], rz[:, j]) for j in range(z.shape[1])]))
    return row


def record_sampling(ctx, checks, model, style, stage, source_beer=None, error=''):
    row = dict(phase=ctx.phase, model=model, style=style, stage=stage, source_beer=source_beer,
               candidate_count=len(checks), sampling_error=error)
    gates = ['chemical_ok', 'chemical_defects_ok', 'stale_proxy_ok', 'flavor_ok',
             'style_supported', 'style_match', 'migration_improved', 'eligible']
    for gate in gates:
        row[gate+'_count'] = int(checks[gate].sum()) if gate in checks else 0
        row[gate+'_rate'] = float(checks[gate].mean()) if len(checks) and gate in checks else np.nan
    for field in DEFECTS:
        key = field+'_pass'
        row[key+'_rate'] = float(checks[key].mean()) if len(checks) and key in checks else np.nan
    if len(checks):
        row['chemical_distance_p50'] = float(checks.chemical_distance.median())
        row['chemical_distance_limit'] = float(checks.chemical_distance_limit.iloc[0])
        row['neighbor_distance_limit'] = float(checks.nearest_distance_limit.iloc[0])
    ctx.sampling_diagnostics.append(row)


def generate_styles(ctx, reference_ids, args):
    items, metrics = [], []
    requested = args.validation_samples if ctx.phase == 'validation' else args.samples
    budget = requested*args.pool_factor
    for si, style in enumerate(ctx.data.requested):
        ref_ids = reference_ids[ctx.data.styles[reference_ids] == style]
        rng = np.random.default_rng(ctx.seed+1000+si*100)
        raw_x = ctx.copula.sample(style, budget, rng)
        raw_checks, raw_pred = ctx.assess(raw_x, style)
        # 两条流程共享同一未约束候选池，避免把随机抽样差异解释为模型改进。
        for name in MODELS:
            note(f'  {ctx.phase}: {STYLE_ZH.get(style, style)} / {model_zh(name)}')
            metrics.append(evaluate_pool(ctx, raw_x, raw_pred, raw_checks, style, name, 'before_screen',
                ctx.data.x[ref_ids], ctx.data.y[ref_ids], requested))
            if name == MODELS[0]:
                x, pred, checks = raw_x, raw_pred, raw_checks
                chosen = np.arange(min(requested, len(x)))
                record_sampling(ctx, checks, name, style, 'unconstrained')
            else:
                pools, predictions, diagnostics = [], [], []
                for round_no in range(args.sampling_rounds):
                    try:
                        proposed = ctx.copula.sample(style, budget, rng, defect_ids=ctx.defect_ids,
                            caps=ctx.defect_caps(style), sweeps=args.gibbs_sweeps)
                    except SamplingInfeasible as exc:
                        record_sampling(ctx, raw_checks.iloc[:0], name, style, 'conditional_sampling', error=str(exc))
                        break
                    table, prediction = ctx.assess(proposed, style)
                    record_sampling(ctx, table, name, style, f'conditional_round_{round_no+1}')
                    pools.append(proposed)
                    predictions.append(prediction)
                    diagnostics.append(table)
                    if sum(int(v.eligible.sum()) for v in diagnostics) >= requested:
                        break
                if pools:
                    x, pred = np.concatenate(pools), np.concatenate(predictions)
                    checks = pd.concat(diagnostics, ignore_index=True)
                else:
                    x, pred, checks = raw_x[:0], raw_pred[:0], raw_checks.iloc[:0].copy()
                metrics.append(evaluate_pool(ctx, x, pred, checks, style, name, 'conditional_candidates',
                    ctx.data.x[ref_ids], ctx.data.y[ref_ids], requested))
                chosen = diversify(x, np.flatnonzero(checks.eligible), checks.screen_score.to_numpy(), requested, ctx.copula.map)
            retained = checks.iloc[chosen].reset_index(drop=True)
            metrics.append(evaluate_pool(ctx, x[chosen], pred[chosen], retained, style, name, 'delivered',
                ctx.data.x[ref_ids], ctx.data.y[ref_ids], requested))
            if len(chosen) < requested:
                note(f'    合格且不重复样本只有 {len(chosen)}/{requested}，不放宽约束、不填充失败样本。')
            items.append(pack(ctx, x[chosen], pred[chosen], retained, name, style))
    return combine(items), pd.DataFrame(metrics)


def migrate(ctx, source_ids, args):
    items, original_items, records = [], [], []
    source_ids = source_ids[ctx.data.styles[source_ids] == 'Pils/Lager']
    target_style = 'Tripel'
    for idx in source_ids:
        source, beer = ctx.data.x[idx], int(ctx.data.meta.beer.iloc[idx])
        original_checks, original_pred = ctx.assess(source[None], target_style, source)
        original_items.append(pack(ctx, source[None], original_pred, original_checks, 'Original', 'Pils/Lager', beer, 'original'))
        for mi, name in enumerate(MODELS):
            note(f'  {ctx.phase}: 啤酒{beer} 皮尔森／拉格 → 三料啤酒 / {model_zh(name)}')
            rng = np.random.default_rng(ctx.seed+50000+beer*10+mi)
            sampler_error = ''
            try:
                x = ctx.copula.sample(target_style, args.migration_pool, rng, source, 'Pils/Lager',
                    defect_ids=ctx.defect_ids, caps=ctx.defect_caps(target_style, source) if name == MODELS[1] else None,
                    sweeps=args.gibbs_sweeps)
            except SamplingInfeasible as exc:
                sampler_error = str(exc)
                # 保存一次无约束失败尝试供诊断，不将它冒充受约束成功样本。
                x = ctx.copula.sample(target_style, args.migration_pool, rng, source, 'Pils/Lager')
            checks, pred = ctx.assess(x, target_style, source)
            record_sampling(ctx, checks, name, target_style, 'migration', beer, sampler_error)
            feasible = np.flatnonzero(checks.eligible.to_numpy()) if not sampler_error else np.array([], dtype=int)
            success = len(feasible)>0
            if success:
                chosen = int(feasible[np.argmin(checks.screen_score.to_numpy()[feasible])])
            else:
                # 失败样本也导出，明确标记，不能算作迁移成功。
                violation = (~checks.constraint_pass).astype(float)*100+(~checks.migration_improved).astype(float)*10
                chosen = int(np.argmin(checks.screen_score.to_numpy()+violation.to_numpy()))
            row = checks.iloc[chosen].to_dict()
            row.update(phase=ctx.phase, model=name, source_beer=beer, source_style='Pils/Lager', target_style=target_style,
                       successful=success, candidate_count=len(x), feasible_count=len(feasible),
                       sampling_error=sampler_error,
                       failure_reason='' if success else 'No candidate met chemistry, defect, flavor, target-style and improvement requirements')
            row['flavor_distance_reduction'] = row['source_flavor_distance']-row['flavor_distance_train_reference']
            row['style_probability_gain'] = row['target_style_probability']-row['source_target_probability']
            row['chemical_distance_reduction'] = row['source_chemical_distance']-row['chemical_distance']
            row['source_audit_flavor_distance'] = float(original_checks.audit_flavor_distance.iloc[0])
            row['audit_flavor_distance_reduction'] = row['source_audit_flavor_distance']-row['audit_flavor_distance']
            row['audit_flavor_improved'] = row['audit_flavor_distance_reduction'] > EPS
            for j, field in enumerate(DEFECTS):
                row[field+'_before'] = source[ctx.defect_ids[j]]
                row[field+'_excess_before'] = max(source[ctx.defect_ids[j]]/ctx.data.thresholds[j]-1, 0)
            row['stale_proxy_before'] = float(original_pred[0, ctx.stale])
            records.append(row)
            items.append(pack(ctx, x[[chosen]], pred[[chosen]], checks.iloc[[chosen]], name, target_style, beer,
                              'successful_migration' if success else 'failed_attempt'))
    if not items:
        raise ValueError('当前迁移阶段没有Pils/Lager源样本。')
    return combine(original_items+items), pd.DataFrame(records)


def summarize_migration(evaluation):
    """成功率分母包括全部源样本；改善量同时报告全部尝试和成功子集。"""
    rows = []
    for (phase, model), frame in evaluation.groupby(['phase', 'model'], sort=False):
        good = frame.loc[frame.successful]
        if frame.source_beer.duplicated().any():
            raise ValueError('同一阶段、模型的迁移源编号重复，成功率分母不可靠。')
        row = dict(phase=phase, model=model, source_count=len(frame), successful_count=len(good),
            failed_count=len(frame)-len(good), success_rate=float(frame.successful.mean()),
            reference_role='heldout_sources' if phase == 'validation' else 'training_sources_descriptive',
            success_and_audit_improved_rate=float((frame.successful & frame.audit_flavor_improved).mean()))
        for label, subset in [('all_attempts', frame), ('successful', good)]:
            for field in ['flavor_distance_reduction', 'chemical_distance_reduction',
                          'style_probability_gain', 'audit_flavor_distance_reduction',
                          'modification_normal_score', 'chemical_defects_ok', 'stale_proxy_ok',
                          'all_known_concentrations_below_threshold', 'style_match']:
                row[label+'_'+field+'_mean'] = float(subset[field].mean()) if len(subset) else np.nan
        rows.append(row)
    return pd.DataFrame(rows)


def grade_reference(ctx, ids):
    meta = ctx.data.meta.iloc[ids][['beer', 'tasting_category_fine']].reset_index(drop=True)
    meta = meta.rename(columns={'beer': 'sample_id', 'tasting_category_fine': 'style'})
    meta['model'] = 'Real_observed'
    grade = q1.quantile_grades(ctx.data.y[ids], ctx.forward.cuts)
    return pd.concat([meta, pd.DataFrame(grade, columns=ctx.data.targets)], axis=1)


def radar_values(grades, groups):
    rows = []
    for (style, model), frame in grades.groupby(['style', 'model'], sort=False):
        for (domain, axis), grouping in groups.groupby(['domain', 'axis'], sort=False):
            values = frame[grouping.target.tolist()].mean(axis=1).to_numpy()
            rows.append(dict(style=style, model=model, domain=domain, axis=axis, label=grouping.label.iloc[0],
                members=grouping.members.iloc[0], n=len(values), mean=float(np.mean(values)),
                p25=float(np.quantile(values, .25)), p75=float(np.quantile(values, .75))))
    return pd.DataFrame(rows)


def explanation(data, args):
    rows = [
        ('生成目标', '只学习给定风格的化合物分布；感官标签不进入生成器。'),
        ('模型对照', '同一个收缩高斯联结模型的原始抽样与受约束抽样对照。已移除条件变分模型。'),
        ('原始基线', '原始流程交付未筛选样本并逐项标记违规；交付不表示合格。受约束流程只交付合格且不重复样本。'),
        ('化学范围修正', '使用训练内部交叉拟合样本校准联合距离、近邻距离、目标风格负对数概率；并保留高斯理论半径。'),
        ('校准范围', '全风格合并校准。样本小，不声称严格的条件覆盖率或共形推断保证。'),
        ('缺陷抽样', f'缺陷边际以{args.gibbs_sweeps}步Gibbs近似截断联合抽样，再条件生成其余化合物；不是逐列裁剪。'),
        ('生成缺陷策略', 'below-threshold；四种化学缺陷均须不高于各自感官阈值。'
         if args.defect_policy == 'below-threshold' else
         f'style-reference；风格参照上限=max(感官阈值,训练浓度{args.defect_quantile:.0%}分位点)，不等于全部低于感官阈值。'),
        ('迁移缺陷约束', '每种化合物的超阈值程度均不得增加；源浓度未知不能认证通过；不得用降低一种缺陷抵消另一种增加。'),
        ('迁移搜索', '从目标分布提议候选，在源谱与目标候选之间搜索；最终仍检查全部约束，不声称全局最优。'),
        ('迁移评价', '成功率以全部源样本为分母；改善量分别汇总全部尝试与成功子集。验证工作簿保留留出源样本的迁移谱和逐项检查。'),
        ('迁移审计', '另报未参与筛选的Ridge感官模型前后距离变化；审计器与主模型共享训练数据，不能代替真实品评。'),
        ('风格识别', '最高概率类别匹配率单列报告，不再作为一票否决条件；筛选使用训练内部折外概率校准的支持门槛。'),
        ('感官预测', f'沿用第一问{data.family}方法重新拟合；模型评价用真实留出感官标签。'),
        ('陈旧酒花', '只检查感官预测代理不增加，并非化学浓度阈值或真实品评保证。'),
        ('第一问评分', '每个维度先按训练七等频边界打0—6整数分；相关组综合分是组内均值，可有小数。'),
        ('雷达图', '50项按香气、风味、口感、其他香气、缺陷分开；仅在同类内用表5正相关合并，每类至少保留三个轴。阴影为样本间四分位范围，不是置信区间。'),
        ('画像与筛选范围', '雷达展示全部50项；现有风味距离仍按30项A_/F_指标计算。缺陷另按逐项约束检查，扩展画像不改变筛选口径。'),
        ('化学字段表示', '四种缺陷列是表8校准浓度，其余保持表1原始表示，可能是对数值，见字段说明。'),
        ('汇总与内标', '五项汇总是近邻估计附加值，不作为生成或约束变量；内标固定训练中位数。'),
        ('验证范围', '先开发留出验证，再用全样本重训最终交付；本次修改参考旧验证结果，不能称为独立最终测试。'),
        ('候选数量', '记录每轮预算、每个筛选门槛通过率与交付覆盖；不足不放宽缺陷约束凑数。'),
        ('参考与阈值', 'trans-2-nonenal阈值来源见阈值表，可按文献范围作敏感性分析；感官阈值不是安全标准。'),
        ('图片标签', '图标题、分面标题及统计说明为中文；化合物使用原表英文名称，感官指标使用原始英文字段，风格保留英文名称。'),
        ('精简运行', str(args.quick)),
    ]
    return pd.DataFrame(rows, columns=['项目', '说明'])


def export_phase(ctx, generated, evaluation, migration, migration_eval, reference_ids, out, args, extra=None):
    data = ctx.data
    migration_summary = summarize_migration(migration_eval)
    real = grade_reference(ctx, reference_ids)
    radar = radar_values(pd.concat([real, generated['grades']], ignore_index=True), data.groups)
    cuts = pd.DataFrame(ctx.forward.cuts.T, columns=[f'q{i}/7' for i in range(1, 7)])
    cuts.insert(0, 'target', data.targets)
    delivered = evaluation.loc[evaluation.stage.eq('delivered')]
    common_styles = set.intersection(*[set(delivered.loc[delivered.model.eq(m) & delivered.n_generated.gt(0), 'style']) for m in MODELS])
    summary_rows = []
    for model in MODELS:
        group = delivered.loc[delivered.model.eq(model)]
        common = group.loc[group['style'].isin(common_styles)]
        summary_rows.append(dict(model=model, requested_styles=len(group),
            styles_with_delivery=int(group.n_generated.gt(0).sum()),
            style_coverage=float(group.n_generated.gt(0).mean()),
            requested_samples=int(group.requested.sum()), generated_samples=int(group.n_generated.sum()),
            delivery_fraction=float(group.n_generated.sum()/group.requested.sum()),
            common_style_count=len(common_styles), common_styles=' | '.join(sorted(common_styles)),
            common_style_profile_MAE_standardized=common.profile_MAE_standardized.mean() if len(common) else np.nan,
            common_style_constraint_pass_rate=common.constraint_pass_rate.mean() if len(common) else np.nan,
            common_style_flavor_pass_rate=common.flavor_pass_rate.mean() if len(common) else np.nan,
            common_style_match_rate=common.target_style_match_rate.mean() if len(common) else np.nan,
            common_style_all_below_threshold_rate=common.all_below_threshold_rate.mean() if len(common) else np.nan,
            common_style_audit_flavor_distance=common.audit_flavor_distance.mean() if len(common) else np.nan))
    tables = {'说明': explanation(data, args), '生成化合物谱': generated['spectra'],
        '感官连续预测': generated['continuous'], '感官0至6评分': generated['grades'],
        '逐样本约束': generated['checks'], '生成评价': evaluation, '模型汇总': pd.DataFrame(summary_rows),
        '雷达图数据': radar, '雷达分组': data.groups, '真实感官评分参照': real,
        '风格感官训练参照': ctx.ref, '评分分位点': cuts, '缺陷阈值': data.threshold_table,
        '化学字段说明': data.catalog, '校准字段核对': data.calibration_changes,
        '感官模型参数': ctx.forward.parameters, '融合权重': ctx.forward.weight_records,
        '范围校准样本': ctx.calibration, '范围校准门槛': ctx.calibration_limits,
        '采样筛选诊断': pd.DataFrame(ctx.sampling_diagnostics), '迁移评价': migration_eval,
        '迁移汇总': migration_summary, '迁移化合物谱': migration['spectra'],
        '迁移感官连续预测': migration['continuous'], '迁移感官0至6评分': migration['grades'],
        '迁移逐项检查': migration['checks']}
    if extra:
        tables.update(extra)
    filename = 'Q3_验证结果.xlsx' if ctx.phase == 'validation' else 'Q3_全部结果.xlsx'
    save_book(out/filename, tables)
    if ctx.phase == 'final':
        for style in FOCUS:
            focus = {k: v for k, v in tables.items() if k in ['说明', '雷达分组', '评分分位点', '缺陷阈值', '化学字段说明']}
            for key in ['生成化合物谱', '感官连续预测', '感官0至6评分', '逐样本约束', '生成评价', '雷达图数据', '真实感官评分参照']:
                focus[key] = tables[key].loc[tables[key]['style'].eq(style)].copy()
            save_book(out/(style+'.xlsx'), focus)
        target_real = real.loc[real['style'].eq('Tripel')]
        paired_radar = []
        for model in MODELS:
            after = migration['grades'].loc[lambda a: a.model.eq(model) & a.status.eq('successful_migration')].copy()
            before = migration['grades'].loc[lambda a: a.model.eq('Original') & a.source_beer.isin(after.source_beer)].copy()
            before['model'] = model+'_before'
            # 每个模型用同一批成功源样本进行前后比较；不混用全部源样本均值。
            frame = radar_values(pd.concat([before, after, target_real], ignore_index=True), data.groups)
            frame['comparison_model'] = model
            paired_radar.append(frame)
        migration_radar = pd.concat(paired_radar, ignore_index=True)
        source_truth_ids = np.array([i for i in ctx.ids if data.styles[i]=='Pils/Lager'])
        actual = pd.DataFrame(data.y[source_truth_ids], columns=data.targets)
        actual.insert(0, 'beer', data.meta.beer.iloc[source_truth_ids].to_numpy())
        mt = {'说明': tables['说明'], '迁移化合物谱': migration['spectra'],
            '成功迁移化合物谱': migration['spectra'].loc[lambda a: a.status.eq('successful_migration')].copy(),
            '迁移感官连续预测': migration['continuous'], '迁移感官0至6评分': migration['grades'],
            '源样本真实感官': actual, '迁移逐项检查': migration['checks'], '迁移评价': migration_eval,
            '迁移汇总': migration_summary,
            '雷达图数据': migration_radar, '雷达分组': data.groups, '缺陷阈值': data.threshold_table,
            '化学字段说明': data.catalog, '评分分位点': cuts}
        save_book(out/'Pils_Lager_to_Tripel.xlsx', mt)
        # 指纹图源数据：固定取各风格均值差异最大的12项设计变量。
        style_means = np.stack([ctx.copula.train_z[ctx.styles == s].mean(0) for s in FOCUS])
        display = np.argsort(np.var(style_means, axis=0))[-12:][::-1]
        rows = []
        for style in FOCUS:
            ref_mean = ctx.copula.train_z[ctx.styles == style].mean(0)
            for j in display:
                rows.append(dict(style=style, model='Real', feature=data.names[j], value=ref_mean[j]))
            for model in MODELS:
                f = generated['spectra'].loc[lambda a: a['style'].eq(style) & a.model.eq(model)]
                if len(f):
                    mean = ctx.copula.map.transform(f[data.names].to_numpy()).mean(0)
                    rows.extend(dict(style=style, model=model, feature=data.names[j], value=mean[j]) for j in display)
        changes = []
        for model in MODELS:
            f = migration['spectra'].loc[lambda a: a.model.eq(model)]
            for _, r in f.iterrows():
                original = migration['spectra'].loc[lambda a: a.model.eq('Original') & a.source_beer.eq(r.source_beer)].iloc[0]
                pair = ctx.copula.map.transform(np.stack([original[data.names].to_numpy(float), r[data.names].to_numpy(float)]))
                for j in range(len(data.names)):
                    changes.append(dict(model=model, source_beer=r.source_beer, status=r.status,
                                        feature=data.names[j], delta_normal_score=pair[1, j]-pair[0, j]))
        save_book(out/'绘图源数据.xlsx', {'重点风格指纹': pd.DataFrame(rows), '迁移化合物变化': pd.DataFrame(changes),
                                       '生成评价': evaluation, '迁移评价': migration_eval})


def plot_results(root):
    """只读工作簿重绘，兼容旧版比较结果；训练逻辑不会从这里触发。"""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.ticker import PercentFormatter

    root = Path(root)
    known = {f.name for f in font_manager.fontManager.ttflist}
    preferred = next((f for f in ['Microsoft YaHei', 'SimHei', 'Noto Sans CJK SC',
                                  'Source Han Sans SC', 'SimSun', 'PingFang SC'] if f in known), None)
    if preferred is None:
        raise RuntimeError('未找到中文字体。请安装微软雅黑、黑体或思源黑体，再用 --plot-only 重画。')
    plt.rcParams.update({'font.family': preferred, 'axes.unicode_minus': False,
                         'pdf.fonttype': 42, 'font.size': 10})
    folder = root/'figures_中文'
    folder.mkdir(exist_ok=True)
    evaluation = pd.read_excel(root/'Q3_验证结果.xlsx', sheet_name='生成评价')
    models = list(dict.fromkeys(evaluation.model))
    colors = {'Real_observed': '#555555', 'Real': '#555555', 'Original': '#777777',
              'GaussianCopula': '#277DA8', 'Copula_constrained': '#D77732',
              'CVAE_screened': '#D77732'}
    colors.update({m+'_before': '#777777' for m in models})
    catalog = pd.read_excel(root/'Q3_全部结果.xlsx', sheet_name='化学字段说明')
    chemical_labels = {r.feature: chemical_display_name(r.chemical_name, r.feature)
                       for r in catalog.itertuples()}
    # 旧版工作簿不回写；仅更新既有图名对照表。
    label_table = catalog[['feature', 'chemical_name']].copy()
    label_table.insert(0, '图中名称', label_table.feature.map(chemical_labels))
    label_table.rename(columns={'feature': '化合物字段', 'chemical_name': '原表化合物名称'}).to_csv(
        folder/'化合物图名对照.csv', index=False, encoding='utf-8-sig')

    def save(fig, name):
        fig.savefig(folder/(name+'.png'), dpi=220, bbox_inches='tight', facecolor='white')
        fig.savefig(folder/(name+'.pdf'), bbox_inches='tight', facecolor='white')
        plt.close(fig)

    def tick_name(model):
        return model_zh(model).replace('高斯联结', '\n高斯联结').replace('条件变分', '\n条件变分')

    def radar(frame, title, filename, extra_note=''):
        panels = []
        for domain in SENSORY_DOMAINS:
            ids = sorted(frame.loc[frame.domain.eq(domain), 'axis'].unique())
            if ids:
                parts = np.array_split(ids, math.ceil(len(ids)/10))
                panels.extend((domain, part, page+1, len(parts)) for page, part in enumerate(parts))
        if not panels:
            raise ValueError('雷达图没有可用的感官分组数据。')
        rows = math.ceil(len(panels)/2)
        height = 6.3*rows+2.
        fig, axes = plt.subplots(rows, 2, subplot_kw={'projection': 'polar'},
                                 figsize=(15, height), squeeze=False)
        handles = {}
        for ax, (domain, part, page, pages) in zip(axes.flat, panels):
            sub = frame.loc[frame.domain.eq(domain) & frame.axis.isin(part)]
            labels = sub[['axis', 'members']].drop_duplicates().sort_values('axis')
            if len(labels) < 3:
                ax.set_axis_off()
                ax.text(.5, .5, '可用感官维度不足，无法绘制雷达图',
                        transform=ax.transAxes, ha='center')
                continue
            angles = np.linspace(0, 2*np.pi, len(labels), endpoint=False)
            closed = np.r_[angles, angles[0]]
            ax.set_theta_offset(np.pi/2)
            ax.set_theta_direction(-1)
            for (style, model), group in sub.groupby(['style', 'model'], sort=False):
                group = group.set_index('axis').reindex(labels.axis)
                values = group['mean'].to_numpy()
                color = colors.get(model, '#777777')
                line = '--' if model in ['Original', 'Real_observed'] or model.endswith('_before') else '-'
                label = f'{style} · {model_zh(model)}（{int(group.n.iloc[0])}个样本）'
                artist, = ax.plot(closed, np.r_[values, values[0]], line, color=color,
                                  linewidth=2, label=label)
                handles[label] = artist
                lo, hi = group.p25.to_numpy(), group.p75.to_numpy()
                ax.fill_between(closed, np.r_[lo, lo[0]], np.r_[hi, hi[0]], color=color, alpha=.09)
            ax.set_xticks(angles)
            ax.set_xticklabels([sensory_group_label(v).replace(' / ', '\n') for v in labels.members],
                               fontsize=9)
            ax.tick_params(axis='x', pad=14)
            ax.set_ylim(0, 6)
            ax.set_yticks([0, 2, 4, 6])
            ax.set_yticklabels(['0', '2', '4', '6'], color='#777777', fontsize=8)
            ax.set_title(SENSORY_DOMAINS[domain][0]+(f'（第{page}部分）' if pages > 1 else ''),
                         pad=48, fontsize=13)
        for ax in list(axes.flat)[len(panels):]:
            ax.set_axis_off()
        fig.suptitle(title, y=1-.2/height, fontsize=17, fontweight='bold')
        fig.text(.5, 1-.75/height, '单项按分位数打０—６分，相关维度取平均；阴影表示样本间四分位范围',
                 ha='center', color='#555555')
        if handles:
            fig.legend(handles.values(), handles.keys(), loc='lower center',
                       bbox_to_anchor=(.5, .55/height), frameon=False, ncol=1, fontsize=10)
        if extra_note:
            fig.text(.5, .2/height, extra_note, ha='center', color='#9B4B24', fontsize=10)
        fig.subplots_adjust(left=.11, right=.89, top=1-1.8/height, bottom=1.8/height,
                            wspace=.58, hspace=.85)
        save(fig, filename)

    fingerprints = pd.read_excel(root/'绘图源数据.xlsx', sheet_name='重点风格指纹')
    for style in FOCUS:
        frame = pd.read_excel(root/(style+'.xlsx'), sheet_name='雷达图数据')
        missing = [model_zh(m) for m in models if m not in set(frame.model)]
        extra = '；'.join(m+'未交付合格样本' for m in missing)
        radar(frame, style+' 感官画像', f'{style}_01_sensory_radar', extra)
        f = fingerprints.loc[fingerprints['style'].eq(style)]
        order = list(dict.fromkeys(f.feature))
        pivot = f.pivot(index='model', columns='feature', values='value').reindex(
            index=['Real', *models], columns=order)
        fig, ax = plt.subplots(figsize=(13, 4.5))
        cmap = plt.get_cmap('RdBu_r').copy()
        cmap.set_bad('#eeeeee')
        im = ax.imshow(np.ma.masked_invalid(pivot.to_numpy()), aspect='auto', cmap=cmap,
                       vmin=-2.5, vmax=2.5)
        ax.set_yticks(range(len(pivot)), [model_zh(m) for m in pivot.index])
        ax.set_xticks(range(len(order)), [textwrap.fill(chemical_labels[v], 22, break_long_words=False,
                      break_on_hyphens=False) for v in order], rotation=38, ha='right', fontsize=9)
        for i in range(len(pivot)):
            if pivot.iloc[i].isna().all():
                ax.text((len(order)-1)/2, i, '无合格生成样本', ha='center', color='#9B4B24')
        ax.set_title(style+' 的化合物特征图', pad=15, fontsize=15)
        fig.colorbar(im, ax=ax, label='训练边际变换后的平均正态分数')
        fig.tight_layout()
        save(fig, f'{style}_02_chemical_fingerprint')

    migration = pd.read_excel(root/'Pils_Lager_to_Tripel.xlsx', sheet_name='迁移评价')
    mradar = pd.read_excel(root/'Pils_Lager_to_Tripel.xlsx', sheet_name='雷达图数据')
    for model in models:
        sub = mradar.loc[mradar.comparison_model.eq(model)]
        f = migration.loc[migration.model.eq(model)]
        count = int(f.successful.sum())
        extra = (f'成功迁移{count}/{len(f)}个；前后曲线仅比较同一批成功样本。' if count else
                 f'成功迁移０/{len(f)}个，仅展示真实目标风格参照；失败候选保留在工作簿中。')
        radar(sub, 'Pils/Lager → Tripel：'+model_zh(model),
              'Migration_01_sensory_radar_'+model, extra)
    delta = pd.read_excel(root/'绘图源数据.xlsx', sheet_name='迁移化合物变化')
    successful = delta.loc[delta.status.eq('successful_migration')]
    fig, ax = plt.subplots(figsize=(13, max(4.8, successful.source_beer.nunique()*.45)))
    if successful.empty:
        ax.axis('off')
        ax.text(.5, .5, '当前约束下没有成功迁移样本。\n失败候选已保留在工作簿中。',
                transform=ax.transAxes, ha='center')
    else:
        order = successful.groupby('feature').delta_normal_score.apply(
            lambda s: np.mean(np.abs(s))).nlargest(12).index
        pivot = successful.loc[successful.feature.isin(order)].pivot(
            index=['model', 'source_beer'], columns='feature', values='delta_normal_score').reindex(columns=order)
        lim = max(.1, np.nanmax(np.abs(pivot.to_numpy())))
        im = ax.imshow(pivot, cmap='RdBu_r', aspect='auto', vmin=-lim, vmax=lim)
        ax.set_xticks(range(len(order)), [textwrap.fill(chemical_labels[v], 22, break_long_words=False,
                      break_on_hyphens=False) for v in order], rotation=38, ha='right', fontsize=9)
        ax.set_yticks(range(len(pivot)), [f'{model_zh(m)}／样品{int(b)}' for m, b in pivot.index], fontsize=9)
        fig.colorbar(im, ax=ax, label='迁移后减迁移前（正态分数）')
    ax.set_title('成功迁移样本的主要化合物变化', pad=15, fontsize=15)
    fig.tight_layout()
    save(fig, 'Migration_02_chemical_changes')

    fig, axes = plt.subplots(1, 4, figsize=(17, 5.2))
    for j, field in enumerate(DEFECTS):
        ax = axes[j]
        for mi, model in enumerate(models):
            f = migration.loc[migration.model.eq(model) & migration.successful.eq(True)]
            if f.empty:
                ax.text(mi*3+.5, .15, '无成功样本', transform=ax.get_xaxis_transform(),
                        ha='center', fontsize=9, color='#9B4B24')
                continue
            before = np.log1p(f[field+'_excess_before'].to_numpy())
            after = np.log1p(f[field+'_excess'].to_numpy())
            for b, a in zip(before, after):
                ax.plot([mi*3, mi*3+1], [b, a], color=colors[model], alpha=.45)
            ax.scatter([mi*3, mi*3+1], [before.mean(), after.mean()], color=colors[model], s=35, zorder=3)
        locations = [mi*3+k for mi in range(len(models)) for k in [0, 1]]
        labels = [tick_name(model)+'\n迁移前' if k == 0 else '迁移后'
                  for model in models for k in [0, 1]]
        ax.set_xticks(locations, labels, fontsize=8)
        ax.set_xlim(-.5, (len(models)-1)*3+1.5)
        ax.set_ylim(bottom=-.02)
        ax.set_title(chemical_labels[field], fontsize=11, pad=15)
        ax.set_ylabel('相对超阈值程度加１的自然对数', fontsize=9)
        ax.grid(axis='y', alpha=.2)
    fig.suptitle('缺陷控制：仅展示成功迁移样本的配对变化', fontsize=16, y=1.01)
    fig.tight_layout()
    save(fig, 'Migration_03_defect_constraints')

    def bar_value(ax, index, value, model, rate=False):
        if np.isfinite(value):
            ax.bar(index, value, width=.55, color=colors[model], alpha=.9)
            ax.annotate(f'{value:.1%}' if rate else f'{value:.3f}', (index, value),
                        xytext=(0, 5 if value >= 0 else -5), textcoords='offset points', ha='center',
                        va='bottom' if value >= 0 else 'top', fontsize=10)
        else:
            ax.text(index, .06, '无可用结果', transform=ax.get_xaxis_transform(),
                    ha='center', color='#9B4B24', fontsize=9)

    validation_migration = pd.read_excel(root/'Q3_验证结果.xlsx', sheet_name='迁移评价')
    fig, axes = plt.subplots(2, 4, figsize=(18, 9.5), squeeze=False)
    for ri, (phase_name, phase_data) in enumerate([
            ('开发验证：留出源样本', validation_migration), ('最终方案：全样本重训', migration)]):
        for mi, model in enumerate(models):
            f = phase_data.loc[phase_data.model.eq(model)]
            audit_reduction = (f.audit_flavor_distance_reduction.mean()
                               if 'audit_flavor_distance_reduction' in f else np.nan)
            values = [f.successful.mean(), f.flavor_distance_reduction.mean(),
                      audit_reduction, f.modification_normal_score.mean()]
            for j, value in enumerate(values):
                bar_value(axes[ri, j], mi, value, model, rate=j == 0)
        for j, title in enumerate(['迁移成功率', '主模型感官距离减少量',
                                   '审计模型感官距离减少量', '化合物谱修改幅度']):
            ax = axes[ri, j]
            ax.set_title(phase_name+'\n'+title, fontsize=11, pad=15)
            ax.set_xticks(range(len(models)), [tick_name(m) for m in models], fontsize=9)
            ax.set_xlim(-.6, len(models)-.4)
            ax.grid(axis='y', alpha=.2)
            ax.margins(y=.22)
            if j in (1, 2):
                ax.axhline(0, color='#777777', linewidth=.8)
        axes[ri, 0].set_ylim(0, 1.15)
        axes[ri, 0].yaxis.set_major_formatter(PercentFormatter(1))
    fig.suptitle('跨风格迁移效果：Pils/Lager → Tripel', fontsize=16)
    fig.text(.5, .02, '均以各阶段全部源样本统计，包含失败尝试；距离减少量为正表示更接近目标。'
             '最终方案属于训练样本描述，不能代替留出验证。', ha='center', fontsize=10)
    fig.tight_layout(rect=(0, .06, 1, .95), h_pad=3)
    save(fig, 'Migration_04_effectiveness')

    delivered = evaluation.loc[evaluation.stage.eq('delivered')]
    common_styles = set.intersection(*[
        set(delivered.loc[delivered.model.eq(m) & delivered.n_generated.gt(0), 'style']) for m in models])
    fig, axes = plt.subplots(1, 3, figsize=(14, 5.5))
    for mi, model in enumerate(models):
        group = delivered.loc[delivered.model.eq(model)]
        common = group.loc[group['style'].isin(common_styles)]
        values = [group.n_generated.sum()/group.requested.sum(),
                  common.profile_MAE_standardized.mean() if len(common) else np.nan,
                  common.constraint_pass_rate.mean() if len(common) else np.nan]
        for j, value in enumerate(values):
            bar_value(axes[j], mi, value, model, rate=j in (0, 2))
    for j, title in enumerate(['开发验证：样本交付率\n（交付不代表满足约束）',
                               '共同交付风格的感官画像误差\n（标准化平均绝对误差，越低越好）',
                               '共同交付风格的化学及缺陷约束通过率']):
        axes[j].set_title(title, fontsize=11, pad=15)
        axes[j].set_xticks(range(len(models)), [tick_name(m) for m in models], fontsize=9)
        axes[j].set_xlim(-.6, len(models)-.4)
        axes[j].grid(axis='y', alpha=.2)
        axes[j].margins(y=.18)
        if j in (0, 2):
            axes[j].set_ylim(0, 1.15)
            axes[j].yaxis.set_major_formatter(PercentFormatter(1))
    fig.suptitle('生成方案的交付与质量评价', fontsize=16)
    comparison_note = (f'共有{len(common_styles)}种风格可作配对质量比较。' if common_styles else
                       '没有共同交付风格，暂不能进行配对质量比较；缺失指标不作为零误差。')
    fig.text(.5, .025, comparison_note, ha='center', color='#9B4B24')
    fig.tight_layout(rect=(0, .08, 1, .93))
    save(fig, 'Evaluation_01_generation_comparison')

    # 各条件分别通过的比例不等于同时通过；帮助定位具体瓶颈。
    checks = pd.read_excel(root/'Q3_全部结果.xlsx', sheet_name='逐样本约束')
    gates = ['chemical_ok', 'chemical_defects_ok', 'stale_proxy_ok', 'flavor_ok', 'eligible']
    titles = ['化学范围', '四种化学缺陷', '陈旧酒花代理', '感官画像', '全部要求']
    fig, ax = plt.subplots(figsize=(12, 5.4))
    width = .75/max(1, len(models))
    x = np.arange(len(gates))
    for mi, model in enumerate(models):
        frame = checks.loc[checks.model.eq(model)]
        offset = (mi-(len(models)-1)/2)*width
        if frame.empty:
            ax.plot([], [], color=colors[model], label=model_zh(model)+'：没有交付样本')
            continue
        values = [frame[k].mean() for k in gates]
        bars = ax.bar(x+offset, values, width, color=colors[model], label=model_zh(model))
        ax.bar_label(bars, labels=[f'{v:.1%}' for v in values], padding=4, fontsize=9)
    ax.set_xticks(x, titles)
    ax.set_ylim(0, 1.22)
    ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.set_ylabel('已交付样本的通过比例')
    ax.set_title('最终生成样本的约束诊断', fontsize=16, pad=18)
    ax.legend(loc='upper center', bbox_to_anchor=(.5, -.13), frameon=False, ncol=len(models))
    ax.grid(axis='y', alpha=.2)
    fig.tight_layout()
    save(fig, 'Evaluation_02_constraint_diagnostics')
    note(f'图片已保存（中文标题、英文指标名称）：{folder}')


def parse_args(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--data-root', type=Path)
    p.add_argument('--output-root', type=Path)
    p.add_argument('--q1-results', type=Path)
    p.add_argument('--forward-model', choices=['auto', *q1.MODELS], default='auto')
    p.add_argument('--styles', nargs='+')
    p.add_argument('--samples', type=int, default=30)
    p.add_argument('--validation-samples', type=int, default=10)
    p.add_argument('--pool-factor', type=int, default=20)
    p.add_argument('--migration-pool', type=int, default=1000)
    p.add_argument('--validation-fraction', type=float, default=.2)
    p.add_argument('--forward-folds', type=int, default=3)
    p.add_argument('--sampling-rounds', type=int, default=2, help='受约束生成最多抽样轮数，不改变门槛')
    p.add_argument('--gibbs-sweeps', type=int, default=60, help='缺陷截断边际的近似Gibbs迭代次数')
    p.add_argument('--support-coverage', type=float, default=.975, help='训练内部折外支持范围校准分位水平')
    p.add_argument('--mean-shrinkage', type=float, default=5.)
    p.add_argument('--cov-shrinkage', type=float, default=20.)
    p.add_argument('--defect-policy', choices=['style-reference', 'below-threshold'], default='style-reference')
    p.add_argument('--defect-quantile', type=float, default=.75)
    p.add_argument('--nonenal-threshold-mg-l', type=float, default=.0001)
    p.add_argument('--radar-min-corr', type=float, default=.35)
    p.add_argument('--flavor-tolerance', type=float, default=1.25)
    p.add_argument('--migration-cost', type=float, default=.5)
    p.add_argument('--jobs', type=int, default=4)
    p.add_argument('--seed', type=int, default=42)
    p.add_argument('--quick', action='store_true')
    p.add_argument('--no-plots', action='store_true')
    p.add_argument('--plot-only', type=Path)
    args = p.parse_args(argv)
    if args.quick:
        args.samples, args.validation_samples, args.pool_factor, args.migration_pool = 4, 3, 4, 30
        args.forward_folds, args.sampling_rounds, args.gibbs_sweeps = 2, 1, 25
        args.styles = list(FOCUS) if args.styles is None else args.styles
    positive = ['samples', 'validation_samples', 'pool_factor', 'migration_pool',
        'sampling_rounds', 'gibbs_sweeps', 'mean_shrinkage', 'cov_shrinkage',
        'nonenal_threshold_mg_l', 'flavor_tolerance', 'jobs']
    if any(not np.isfinite(getattr(args, k)) or getattr(args, k)<=0 for k in positive):
        p.error('数量、迭代次数、收缩强度、阈值和jobs必须为有限正数。')
    if not (0<args.validation_fraction<.5 and 0<args.defect_quantile<1 and 0<args.radar_min_corr<1
            and .5<args.support_coverage<1):
        p.error('验证比例须在(0,0.5)，分位点和相关阈值须在(0,1)。')
    if (args.forward_folds<2 or not np.isfinite(args.migration_cost) or args.migration_cost<0
            or not 0<=args.seed<2**31-100000):
        p.error('折数、修改成本或随机种子无效。')
    return args


def main(argv=None):
    args = parse_args(argv)
    if args.plot_only:
        importlib.import_module('openpyxl')
        plot_results(args.plot_only.expanduser().resolve())
        return
    check_dependencies(not args.no_plots)
    data = load_data(args)
    parent = args.output_root or data.root/'运行结果'
    out = parent/('Q3_'+('QUICK_' if args.quick else '')+datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    out.mkdir(parents=True, exist_ok=False)
    train, test = stratified_holdout(data.styles, args.validation_fraction, args.seed)
    split_table = data.meta.copy()
    split_table['role'] = np.where(np.isin(np.arange(len(data.x)), train), 'train', 'heldout')
    manifest = dict(status='running', protocol=PROTOCOL, parameters=vars(args),
        forward_family=data.family, q1_result_folder=data.q1_folder,
        source_hashes={str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in data.files.values()},
        script_hash=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        q1_script_hash=hashlib.sha256(Path(q1.__file__).read_bytes()).hexdigest(),
        python=platform.python_version(), trained_models_not_experimental_evidence=True,
        generation_condition='tasting_category_fine only',
        validation_scope='development_holdout; earlier results informed this revision', quick=args.quick)
    write_json(manifest, out/'manifest.json')
    note(f'输出目录：{out}\n感官模型：{data.family}；风格数：{len(data.requested)}')
    try:
        note('阶段1/2：真实留出验证。')
        validation = Context(data, train, args, 'validation', args.seed)
        pred = validation.forward.predict(data.x[test])
        audit = validation.forward.audit(data.x[test])
        forward_metrics = pd.concat([sensory_metrics(data.y[test], pred, data.targets, 'heldout', data.family),
            sensory_metrics(data.y[test], audit, data.targets, 'heldout', 'Ridge_auditor')], ignore_index=True)
        heldout = pd.concat([data.meta.iloc[test].reset_index(drop=True), pd.DataFrame(pred, columns=data.targets)], axis=1)
        style_probability = validation.classifier.predict_proba(validation.copula.map.transform(data.x[test]))
        style_prediction = validation.classifier.classes_[style_probability.argmax(1)]
        style_metrics = pd.DataFrame([{'phase': 'heldout', 'n': len(test),
            'accuracy': accuracy_score(data.styles[test], style_prediction),
            'balanced_accuracy': balanced_accuracy_score(data.styles[test], style_prediction),
            'macro_F1': f1_score(data.styles[test], style_prediction, average='macro', zero_division=0),
            'log_loss': log_loss(data.styles[test], style_probability, labels=validation.classifier.classes_)}])
        heldout_style = data.meta.iloc[test].reset_index(drop=True).copy()
        heldout_style['predicted_style'] = style_prediction
        generated, evaluation = generate_styles(validation, test, args)
        migration, migration_eval = migrate(validation, test, args)
        export_phase(validation, generated, evaluation, migration, migration_eval, test, out, args,
                     {'真实样本划分': split_table, '留出感官预测指标': forward_metrics, '留出感官预测': heldout,
                      '留出风格识别指标': style_metrics, '留出风格识别': heldout_style})
        manifest['status'] = 'validation_complete'
        write_json(manifest, out/'manifest.json')
        del validation
        note('阶段2/2：全部真实样本重新拟合，生成最终交付和迁移方案。')
        final = Context(data, np.arange(len(data.x)), args, 'final', args.seed+10000)
        generated, evaluation = generate_styles(final, np.arange(len(data.x)), args)
        migration, migration_eval = migrate(final, np.arange(len(data.x)), args)
        export_phase(final, generated, evaluation, migration, migration_eval, np.arange(len(data.x)), out, args,
                     {'真实样本划分': split_table, '留出感官预测指标': forward_metrics, '留出风格识别指标': style_metrics})
        manifest['status'] = 'computed'
        write_json(manifest, out/'manifest.json')
        if not args.no_plots:
            plot_results(out)
        manifest['status'] = 'complete'
        write_json(manifest, out/'manifest.json')
        note(f'完成：{out}\n先看验证结果与约束通过率，再看重点风格及迁移工作簿。')
    except BaseException as exc:
        manifest['error'] = f'{type(exc).__name__}: {exc}'
        manifest['status'] = 'computed_plot_failed' if manifest['status']=='computed' else 'failed'
        write_json(manifest, out/'manifest.json')
        raise


if __name__ == '__main__':
    main()
