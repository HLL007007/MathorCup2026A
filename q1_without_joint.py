"""第一问：六种回归方法的交叉验证，以及最佳模型的感官评分输出。

python q1_without_joint.py                 # 正式实验并画三类图
python q1_without_joint.py --quick         # 精简候选参数，仅用于流程检查
python q1_without_joint.py --plot-only     # 重画最近一次完整的本版结果
"""
from __future__ import annotations

import argparse
import hashlib
import json
import platform
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import sklearn
from joblib import Parallel, delayed
from scipy.optimize import minimize
from sklearn.cross_decomposition import PLSRegression
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import Ridge
from sklearn.model_selection import KFold
from sklearn.preprocessing import StandardScaler

BASE_MODELS = ("Ridge", "PLS2", "RF", "GBDT")
FUSION_MODELS = ("TargetSelect", "TargetBlend")
MODELS = BASE_MODELS + FUSION_MODELS
PROTOCOL = "six_models_best_scores_v4"
METRICS = ("MAE", "RMSE", "R2", "RMSE_pooled")
META_COLUMNS = {"beer", "beer_id", "tasting_category_fine"}
PALATE_OVERALL = {"alcohol", "aftertaste", "body", "co2", "overall"}
SUMMARY_ALIASES = {"A_ester_fruity": "A_esters_fruity",
                   "F_ester_fruity": "F_esters_fruity", "4vg": "X4vg"}
GRADE_PROBABILITIES = np.arange(1, 7, dtype=float) / 7
REQUIRED_RESULTS = ("run_manifest.json", "cv_summary.csv", "oof_predictions.csv", "quantile_thresholds.csv", "best_model.json")



def resolve_data_root(provided):
    """优先使用显式路径；否则根据脚本位置和当前工作目录查找项目。"""
    if provided is not None:
        return provided.expanduser().resolve()
    script_dir = Path(__file__).resolve().parent
    candidates = [
        script_dir,
        script_dir.parent,
        Path.cwd(),
        script_dir / "MathorCup2026A",
        Path.cwd() / "MathorCup2026A",
    ]
    checked = []
    for candidate in candidates:
        root = candidate.resolve()
        if root in checked:
            continue
        checked.append(root)
        if (root / "处理结果/最终数据/4_专业感官数据.csv").is_file():
            return root
    locations = "\n".join(str(root) for root in checked)
    raise FileNotFoundError(
        "无法自动定位数据目录。请用 --data-root 指向含有“处理结果”的项目目录。"
        f"\n已检查：\n{locations}"
    )


def read_keyed_csv(path):
    frame = pd.read_csv(path, encoding="utf-8-sig", dtype={"beer_id": "string"})
    if "beer" not in frame:
        raise ValueError(f"缺少 beer 编号列：{path}")
    beer = pd.to_numeric(frame["beer"], errors="raise")
    if beer.isna().any() or (beer != np.floor(beer)).any() or beer.duplicated().any():
        raise ValueError(f"beer 编号必须为无缺失、不重复的整数：{path}")
    frame["beer"] = beer.astype("int64")
    return frame.set_index("beer").sort_index()


def load_arrays(root, target_scope):
    root = root.resolve()
    files = {
        "chemical_raw_missing": root / "处理结果/处理记录/结构整理（第4版）/02_化合物_保留缺失.csv",
        "chemical_final_schema": root / "处理结果/最终数据/1_化合物数据_已填补标准化.csv",
        "sensory": root / "处理结果/最终数据/4_专业感官数据.csv",
    }
    for path in files.values():
        if not path.is_file():
            raise FileNotFoundError(f"找不到文件：{path}\n请用 --data-root 指向 MathorCup2026A 目录。")
    raw = read_keyed_csv(files["chemical_raw_missing"])
    schema = read_keyed_csv(files["chemical_final_schema"])
    sensory = read_keyed_csv(files["sensory"])
    if not raw.index.equals(schema.index) or not raw.index.equals(sensory.index):
        raise ValueError("三张表的 beer 集合不一致，拒绝静默丢弃样本。")
    features = [c for c in schema.columns if c not in META_COLUMNS]
    if len(features) != 231 or set(features) - set(raw.columns):
        raise ValueError("化学字段应为 231 项，且保留缺失副本必须包含全部字段。")
    targets = [c for c in sensory.columns if c not in META_COLUMNS]
    if target_scope == "af":
        targets = [c for c in targets if c.startswith(("A_", "F_"))]
    elif target_scope == "flavor":
        targets = [c for c in targets if c not in PALATE_OVERALL]
    expected = {"all": 50, "af": 30, "flavor": 45}[target_scope]
    if len(targets) != expected:
        raise ValueError(f"当前目标字段数 {len(targets)} 与 {target_scope} 预期 {expected} 不一致。")
    X = raw[features].apply(pd.to_numeric, errors="raise").to_numpy(dtype=float)
    Y = sensory[targets].apply(pd.to_numeric, errors="raise").to_numpy(dtype=float)
    if np.isinf(X).any() or not np.isfinite(Y).all():
        raise ValueError("X 不允许无穷值，Y 不允许缺失或无穷值。")
    for column in ("beer_id", "tasting_category_fine"):
        if not schema[column].astype(str).equals(sensory[column].astype(str)):
            raise ValueError(f"样本信息不一致：{column}")
    meta = sensory[["beer_id", "tasting_category_fine"]].reset_index()
    return X, Y, meta, features, targets, files


def read_summary_columns(path):
    """只读取 XLSX 的 summary 表头；不读取评分公式，不依赖 Excel/openpyxl。"""
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    with zipfile.ZipFile(path) as book:
        workbook = ET.fromstring(book.read("xl/workbook.xml"))
        sheets = [s for s in workbook.findall("s:sheets/s:sheet", ns)
                  if s.get("name", "").casefold() == "summary"]
        if len(sheets) != 1:
            raise ValueError("3 号工作簿必须有且仅有一个 summary 工作表。")
        relations = ET.fromstring(book.read("xl/_rels/workbook.xml.rels"))
        relation = next(r for r in relations if r.get("Id") == sheets[0].get(rel_attr))
        if relation.get("TargetMode") == "External":
            raise ValueError("summary 工作表不能为外部链接。")
        target = relation.get("Target")
        sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        strings = []
        if "xl/sharedStrings.xml" in book.namelist():
            shared = ET.fromstring(book.read("xl/sharedStrings.xml"))
            strings = ["".join(t.text or "" for t in si.findall(".//s:t", ns)) for si in shared]
        root = ET.fromstring(book.read(sheet_path))
        row = root.find("s:sheetData/s:row[@r='1']", ns)
        if row is None:
            raise ValueError("summary 第一行缺少表头。")
        columns = []
        for cell in row.findall("s:c", ns):
            value = cell.find("s:v", ns)
            if cell.get("t") == "s" and value is not None:
                label = strings[int(value.text)]
            elif cell.get("t") == "inlineStr":
                label = "".join(t.text or "" for t in cell.findall(".//s:t", ns))
            else:
                label = value.text if value is not None else ""
            if not label.strip():
                raise ValueError("summary 表头出现空列，请检查评分卡。")
            columns.append(label.strip())
    if len(columns) != len(set(columns)) or columns[0] != "Beer" or columns[-1] != "comments":
        raise ValueError("summary 表头重复或首尾字段与 Beer/comments 不符。")
    return columns



def parse_args(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--data-root", type=Path)
    p.add_argument("--output-root", type=Path)
    p.add_argument("--targets", choices=["all", "af", "flavor"], default="all")
    p.add_argument("--folds", type=int, default=5)
    p.add_argument("--inner-folds", type=int, default=3)
    p.add_argument("--inner-repeats", type=int, default=3,
                   help="内层交叉验证重复次数；默认 3 次，每次覆盖全部外层训练样本")
    p.add_argument("--repeats", type=int, default=1)
    p.add_argument("--seed", type=int, default=42)
    p.add_argument("--jobs", type=int, default=4)
    p.add_argument("--quick", action="store_true")
    p.add_argument("--clip-quantile", type=float, default=0.01,
                   help="训练折缩尾比例；0 表示关闭")
    p.add_argument("--selection-metric", choices=["MAE", "RMSE", "R2"], default="MAE",
                   help="基模型调参及 TargetSelect 的目标，默认 MAE；TargetBlend 仍用平方损失")
    p.add_argument("--blend-shrinkage", type=float, default=0.05,
                   help="融合权重向该目标的内层最佳基模型收缩的强度，预先固定；默认 0.05")
    p.add_argument("--plot-only", nargs="?", const="latest", default=None,
                   help="只画已有本版结果；不指定路径则寻找最近一次完整结果")
    p.add_argument("--plot-repeat", type=int, default=1)
    p.add_argument("--quantile-fold", type=int, default=1)
    p.add_argument("--heatmap-beers", nargs=12, type=int, default=list(range(1, 13)),
                   help="12 个 beer 编号，默认 1–12")
    p.add_argument("--no-plots", action="store_true")
    a = p.parse_args(argv)
    if not 0 <= a.clip_quantile < 0.1:
        p.error("缩尾比例应在 [0,0.1)。")
    if min(a.folds, a.inner_folds) < 2 or min(a.repeats, a.inner_repeats) < 1:
        p.error("内外折数至少为 2，重复次数至少为 1。")
    if a.jobs == 0 or a.jobs < -1 or min(a.plot_repeat, a.quantile_fold) < 1:
        p.error("jobs 应为正整数或 -1；绘图重复编号和折编号至少为 1。")
    if not np.isfinite(a.blend_shrinkage) or a.blend_shrinkage < 0:
        p.error("--blend-shrinkage 必须为有限非负数。")
    if len(set(a.heatmap_beers)) != 12:
        p.error("热力图的 12 个 beer 编号不能重复。")
    if a.no_plots and a.plot_only is not None:
        p.error("--no-plots 与 --plot-only 不能同时使用。")
    if a.plot_only is None and (a.plot_repeat > a.repeats or a.quantile_fold > a.folds):
        p.error("绘图的重复编号或折编号超过本次运行范围。")
    try:
        a.data_root = resolve_data_root(a.data_root)
    except FileNotFoundError as error:
        p.error(str(error))
    return a


@dataclass
class Data:
    X: np.ndarray
    Y: np.ndarray
    meta: pd.DataFrame
    chemicals: list
    targets: list
    summary_columns: list
    files: dict


def load_data(args):
    X, Y, meta, chemicals, targets, files = load_arrays(args.data_root, args.targets)
    files["summary_template"] = args.data_root / "处理结果/最终数据/3_感官评分卡.xlsx"
    columns = read_summary_columns(files["summary_template"])
    all_targets = [c for c in pd.read_csv(files["sensory"], nrows=0).columns if c not in META_COLUMNS]
    mapped = [SUMMARY_ALIASES.get(c, c) for c in columns[1:-1]]
    if len(mapped) != 50 or len(set(mapped)) != 50 or set(mapped) != set(all_targets):
        raise ValueError("summary 的 50 个感官字段不能与表格 4 一一对应。")
    missing = set(args.heatmap_beers) - set(meta.beer)
    if missing:
        raise ValueError(f"热力图编号在数据中不存在：{sorted(missing)}")
    return Data(X, Y, meta, chemicals, targets, columns, files)


class FoldPreprocessor:
    """训练折内拟合中位数填补、缩尾和标准化，保留全部化学字段。"""
    def __init__(self, clip_quantile):
        self.clip_quantile = clip_quantile

    def fit(self, X):
        self.imputer = SimpleImputer(strategy="median", keep_empty_features=True)
        values = self.imputer.fit_transform(X)
        if self.clip_quantile:
            self.lower, self.upper = np.quantile(values, [self.clip_quantile, 1-self.clip_quantile], axis=0)
        else:
            self.lower, self.upper = np.full(values.shape[1], -np.inf), np.full(values.shape[1], np.inf)
        self.scaler = StandardScaler().fit(np.clip(values, self.lower, self.upper))
        return self

    def transform(self, X):
        values = self.imputer.transform(X)
        result = self.scaler.transform(np.clip(values, self.lower, self.upper))
        if not np.isfinite(result).all():
            raise ValueError("预处理特征存在非有限值。")
        return result


def prepare_fold(data, train, valid, args):
    transform = FoldPreprocessor(args.clip_quantile).fit(data.X[train])
    return dict(x_train=transform.transform(data.X[train]), x_valid=transform.transform(data.X[valid]),
                y_train=data.Y[train], y_valid=data.Y[valid], transform=transform)


def predict_model(name, parameters, fold, args, seed):
    x, xv = fold["x_train"], fold["x_valid"]
    scaler = StandardScaler().fit(fold["y_train"])
    y = scaler.transform(fold["y_train"])
    per_target = isinstance(parameters, list)
    if per_target and len(parameters) != y.shape[1]:
        raise ValueError("逐目标参数数量与感官目标数不一致。")
    if name == "Ridge":
        # sklearn 原生多输出 Ridge 支持每个输出独立的 alpha。
        alpha = np.array([item["alpha"] for item in parameters]) if per_target else parameters["alpha"]
        model = Ridge(alpha=alpha, solver="svd")
        prediction = np.asarray(model.fit(x, y).predict(xv))
    elif name == "PLS2":
        if per_target:
            raise ValueError("PLS2 联合拟合全部输出，必须使用统一的潜变量数。")
        components = max(1, min(parameters["n_components"], x.shape[1], len(x)-1))
        model = PLSRegression(n_components=components, scale=False, max_iter=1000, tol=1e-7)
        prediction = np.asarray(model.fit(x, y).predict(xv))
    elif name in {"RF", "GBDT"}:
        def one_target(j):
            selected = parameters[j] if per_target else parameters
            if name == "RF":
                model = RandomForestRegressor(**selected, random_state=seed+j, n_jobs=1)
            else:
                model = GradientBoostingRegressor(**selected, random_state=seed+j)
            return model.fit(x, y[:, j]).predict(xv)
        prediction = np.column_stack(Parallel(n_jobs=args.jobs, prefer="threads")(
            delayed(one_target)(j) for j in range(y.shape[1])))
    else:
        raise ValueError(f"未知模型：{name}")
    prediction = scaler.inverse_transform(prediction.reshape(len(xv), y.shape[1]))
    if not np.isfinite(prediction).all():
        raise ValueError(f"{name} 产生非有限预测。")
    return prediction


def finite_mean(values, axis=0):
    """常数目标的 R2 保留 NaN，避免用 0 冒充有效得分。"""
    values = np.asarray(values, dtype=float)
    count = np.isfinite(values).sum(axis=axis)
    total = np.where(np.isfinite(values), values, 0).sum(axis=axis)
    return np.divide(total, count, out=np.full_like(total, np.nan, dtype=float), where=count > 0)


def select_candidate_indices(metrics, selection_metric, joint=False):
    """metrics 为候选×目标；R2 完全不可定义的目标显式回退 MAE。"""
    values = np.asarray(metrics[selection_metric], dtype=float)
    losses = -values if selection_metric == "R2" else values
    effective = np.full(values.shape[1], selection_metric, dtype=object)
    undefined = ~np.isfinite(losses).any(axis=0)
    losses = np.where(np.isfinite(losses), losses, np.inf)
    if undefined.any():
        losses[:, undefined] = metrics["MAE"][:, undefined]
        effective[undefined] = "MAE_fallback_constant_target"
    if joint:
        # PLS2 的 R2 调参只平均可定义目标；全为常数时才整体回退 MAE。
        use = ~undefined
        aggregate = losses[:, use].mean(axis=1) if use.any() else metrics["MAE"].mean(axis=1)
        indices = np.full(values.shape[1], int(np.argmin(aggregate)), dtype=int)
    else:
        indices = np.argmin(losses, axis=0)
    return indices, effective


def select_model(name, folds, args, seed, repeat, outer_fold, targets, records, params_records):
    candidates = make_candidates(name, args.quick)
    all_predictions, candidate_metrics = [], {key: [] for key in ("MAE", "RMSE", "R2")}
    for number, parameters in enumerate(candidates, 1):
        details, predictions = [], []
        for i, fold in enumerate(folds):
            prediction = predict_model(name, parameters, fold, args, seed+i+1)
            predictions.append(prediction)
            _, detail = metric_values(fold["y_valid"], prediction)
            details.append(detail)
        metrics = {key: finite_mean([item[key] for item in details]) for key in candidate_metrics}
        for key in candidate_metrics:
            candidate_metrics[key].append(metrics[key])
        all_predictions.append(predictions)
        for j, target in enumerate(targets):
            records.append(dict(repeat=repeat, fold=outer_fold, model=name,
                target=target, candidate=number, parameters=json.dumps(parameters),
                **{key: float(value[j]) for key, value in metrics.items()}))
        value = float(finite_mean(metrics[args.selection_metric]))
        print(f"    {name} 候选 {number}: 内层宏平均 {args.selection_metric}={value:.5f}", flush=True)
    metrics = {key: np.asarray(value) for key, value in candidate_metrics.items()}
    indices, effective = select_candidate_indices(metrics, args.selection_metric, joint=(name == "PLS2"))
    selected = [dict(candidates[i]) for i in indices]
    for j, target in enumerate(targets):
        params_records.append(dict(repeat=repeat, fold=outer_fold, model=name,
            target=target, candidate=int(indices[j])+1, parameters=json.dumps(selected[j]),
            tuning_scope="joint_multioutput" if name == "PLS2" else "per_target",
            selection_metric=args.selection_metric, effective_metric=str(effective[j])))
    # 每个目标取自己的获胜候选；PLS2 的全部目标取同一个多输出候选。
    parts = [np.column_stack([all_predictions[indices[j]][i][:, j] for j in range(len(targets))])
             for i in range(len(folds))]
    return (selected[0] if name == "PLS2" else selected), parts


def assemble_inner_oof(predictions, splits, y_train, repeats=1):
    """每次内层重复覆盖各样本一次，然后按样本平均重复预测。"""
    if len(predictions) != len(splits):
        raise ValueError("内层预测个数与划分个数不一致。")
    result = np.zeros_like(y_train, dtype=float)
    coverage = np.zeros(len(y_train), dtype=int)
    for prediction, (train, valid) in zip(predictions, splits):
        if np.intersect1d(train, valid).size or len(np.unique(valid)) != len(valid):
            raise ValueError("内层训练/验证编号重叠或验证编号重复。")
        if np.asarray(prediction).shape != (len(valid), y_train.shape[1]):
            raise ValueError("内层预测的样本数或目标数不符。")
        result[valid] += prediction
        coverage[valid] += 1
    if not np.all(coverage == repeats) or not np.isfinite(result).all():
        raise ValueError("内层折外预测覆盖次数不等于内层重复次数。")
    return result / repeats


def learn_target_fusion(y_train, inner_oof, targets, shrinkage, selection_metric="MAE"):
    """只接收外层训练标签和训练内部的 OOF 预测，不接收外层测试数据。

    TargetSelect 按指定指标选择基模型（默认 MAE）。TargetBlend 最小化：
      mean(((P_j @ w - y_j) / scale_j)**2) + lambda * ||w - one_hot_best||²
    约束 w>=0、sum(w)=1；无截距、不混合不同目标，lambda 不看外层结果调节。
    内层 OOF 的基模型参数已用相同内层折选出；其误差仅供训练诊断。
    """
    y = np.asarray(y_train, dtype=float)
    p = np.asarray(inner_oof, dtype=float)
    if y.ndim != 2 or p.shape != (*y.shape, len(BASE_MODELS)) or len(targets) != y.shape[1]:
        raise ValueError("融合输入应为 n×目标数×4，目标顺序必须一致。")
    if not np.isfinite(y).all() or not np.isfinite(p).all():
        raise ValueError("融合训练标签与预测必须为有限值。")
    if not np.isfinite(shrinkage) or shrinkage < 0:
        raise ValueError("融合收缩强度必须为有限非负数。")
    mse = np.mean((p-y[:, :, None])**2, axis=0)
    maes = np.mean(np.abs(p-y[:, :, None]), axis=0)
    if selection_metric not in {"MAE", "RMSE", "R2"}:
        raise ValueError("未知融合选择指标。")
    # 同一目标的 R2 最大化等价于 MSE 最小化；常数目标回退 MAE。
    loss = maes.copy() if selection_metric == "MAE" else mse.copy()
    constant = np.sum((y-y.mean(axis=0))**2, axis=0) <= 1e-12
    if selection_metric == "R2":
        loss[constant] = maes[constant]
    best = np.argmin(loss, axis=1)  # 相同误差按 BASE_MODELS 的固定顺序决胜。
    hard = np.eye(len(BASE_MODELS))[best]
    weights = np.zeros_like(hard)
    records = []
    for j,target in enumerate(targets):
        center = float(np.mean(y[:,j]))
        scale = float(np.std(y[:,j]))
        # 常数/近常数标签用 1 作缩放，避免把接近 0 的方差作为除数。
        scale = scale if scale > 1e-8 else 1.0
        design = (p[:,j,:]-center)/scale
        response = (y[:,j]-center)/scale
        anchor = hard[j]
        def objective(w):
            error = design @ w-response
            return float(np.mean(error**2)+shrinkage*np.sum((w-anchor)**2))
        def gradient(w):
            return 2*design.T@(design@w-response)/len(y)+2*shrinkage*(w-anchor)
        optimum = minimize(objective, anchor.copy(), jac=gradient, method="SLSQP",
            bounds=[(0.0,1.0)]*len(BASE_MODELS),
            constraints={"type":"eq", "fun":lambda w:float(w.sum()-1.0),
                         "jac":lambda w:np.ones_like(w)},
            options={"maxiter":300,"ftol":1e-10})
        candidate = np.asarray(optimum.x, dtype=float)
        valid = (optimum.success and np.isfinite(candidate).all()
                 and np.min(candidate)>=-1e-7 and abs(candidate.sum()-1)<1e-6)
        status = "optimized"
        if valid:
            candidate = np.clip(candidate,0,1)
            candidate /= candidate.sum()
            # anchor 是可行点，数值求解结果不应比它有更高的正则化目标值。
            valid = objective(candidate) <= objective(anchor)+1e-9
        if not valid:
            candidate = anchor.copy()
            status = "fallback_to_target_select"
        weights[j] = candidate
        blend_prediction = p[:,j,:] @ candidate
        records.append(dict(target=target, selected_model=BASE_MODELS[best[j]],
            selection_metric=selection_metric,
            effective_selection_metric="MAE_fallback_constant_target" if selection_metric == "R2" and constant[j] else selection_metric,
            select_train_oof_MAE=float(maes[j,best[j]]),
            blend_train_oof_MAE=float(np.mean(np.abs(y[:,j]-blend_prediction))),
            select_train_oof_RMSE=float(np.sqrt(mse[j,best[j]])),
            blend_train_oof_RMSE=float(np.sqrt(np.mean((y[:,j]-blend_prediction)**2))),
            blend_shrinkage=float(shrinkage), target_scale=scale,
            optimizer_status=status, optimizer_message=str(optimum.message),
            **{f"oof_RMSE_{name}":float(np.sqrt(mse[j,k])) for k,name in enumerate(BASE_MODELS)},
            **{f"oof_MAE_{name}":float(maes[j,k]) for k,name in enumerate(BASE_MODELS)},
            **{f"weight_{name}":float(candidate[k]) for k,name in enumerate(BASE_MODELS)}))
    return {"TargetSelect":hard, "TargetBlend":weights}, records


def apply_target_fusion(base_predictions, weights):
    """仅组合同一目标的四个连续预测；不读取标签，不做整数分级。"""
    p, w = np.asarray(base_predictions, dtype=float), np.asarray(weights, dtype=float)
    if p.ndim != 3 or w.shape != p.shape[1:] or p.shape[2] != len(BASE_MODELS):
        raise ValueError("融合预测与权重维度不一致。")
    if not np.isfinite(p).all() or not np.isfinite(w).all() or (w<0).any() or not np.allclose(w.sum(axis=1),1,atol=1e-8):
        raise ValueError("融合权重必须非负、逐目标和为 1，预测必须有限。")
    return np.einsum("njm,jm->nj",p,w)


def export_fusion_comparison(folds, output):
    """按重复编号和外层折比较融合与各基模型。"""
    records = []
    for fusion in FUSION_MODELS:
        for base in BASE_MODELS:
            pairs = folds[folds.model == base].merge(folds[folds.model == fusion],
                on=["repeat", "fold"], suffixes=("_base", "_fusion"), validate="one_to_one")
            for _, row in pairs.iterrows():
                item = dict(baseline=base, fusion=fusion, repeat=int(row["repeat"]), fold=int(row["fold"]))
                for metric in ("MAE", "RMSE", "R2"):
                    a, b = row[f"{metric}_base"], row[f"{metric}_fusion"]
                    item[f"{metric}_improvement"] = b-a if metric == "R2" else a-b
                records.append(item)
    frame = pd.DataFrame(records)
    frame.to_csv(output / "paired_fusion_comparison.csv", index=False, encoding="utf-8-sig")
    metrics = [f"{metric}_improvement" for metric in ("MAE", "RMSE", "R2")]
    summary = frame.groupby(["baseline", "fusion"], sort=False)[metrics].agg(["mean", "std"])
    summary.columns = [f"{a}_{b}" for a, b in summary.columns]
    summary.reset_index().to_csv(output / "fusion_comparison_summary.csv", index=False, encoding="utf-8-sig")


def save_csv(rows, path):
    pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8-sig")



def make_candidates(kind, quick):
    if kind == "Ridge":
        return [{"alpha": a} for a in ([100.] if quick else [0.1, 1., 10., 100., 1000., 10000.])]
    if kind == "PLS2":
        return [{"n_components": a} for a in ([5] if quick else [2, 5, 10, 15])]
    if kind == "RF":
        return [dict(n_estimators=60 if quick else 150, max_depth=None, min_samples_leaf=leaf,
                     max_features=features) for leaf in ([5] if quick else [2, 5])
                for features in ([0.5] if quick else [0.5, 0.8])]
    if kind == "GBDT":
        return [dict(n_estimators=60 if quick else 150, learning_rate=0.05, max_depth=depth,
                     min_samples_leaf=5, subsample=1.0, loss="squared_error")
                for depth in ([2] if quick else [1, 2])]
    raise ValueError(f"不是基模型：{kind}")


def metric_values(y_true, y_pred):
    error = y_true - y_pred
    mae = np.mean(np.abs(error), axis=0)
    rmse = np.sqrt(np.mean(error ** 2, axis=0))
    sse = np.sum(error ** 2, axis=0)
    sst = np.sum((y_true - y_true.mean(axis=0)) ** 2, axis=0)
    r2 = np.full(y_true.shape[1], np.nan)
    valid = sst > 1e-12
    r2[valid] = 1.0 - sse[valid] / sst[valid]
    summary = {"MAE": float(mae.mean()), "RMSE": float(rmse.mean()),
               "R2": float(np.mean(r2[valid])) if valid.any() else np.nan,
               "RMSE_pooled": float(np.sqrt(np.mean(error ** 2))),
               "R2_valid_targets": int(valid.sum()), "n_targets": y_true.shape[1]}
    return summary, {"MAE": mae, "RMSE": rmse, "R2": r2}


def export_summary(frame, keys, path):
    result = frame.groupby(keys, sort=False)[list(METRICS)].agg(["mean", "std"])
    result.columns = [f"{metric}_{statistic}" for metric, statistic in result.columns]
    result.reset_index().to_csv(path, index=False, encoding="utf-8-sig")
    return result.reset_index()


def fit_grade_cuts(y_train):
    """每个目标独立计算六个等频分位点，只接受当前训练折的连续标签。"""
    values = np.asarray(y_train, dtype=float)
    if values.ndim != 2 or not values.size or not np.isfinite(values).all():
        raise ValueError("分位点拟合需要非空、有限的二维训练标签。")
    return np.quantile(values, GRADE_PROBABILITIES, axis=0, method="linear")


def quantile_grades(values, cuts):
    """(-inf,q1] 为 0，(q1,q2] 为 1，...，(q6,inf) 为 6。

    重复阈值保留；相同值落入同一档，因此有并列值时不保证七档人数完全相等。
    阈值之外的预测自然落入 0/6 档，不对连续预测本身四舍五入。
    """
    values, cuts = np.asarray(values, dtype=float), np.asarray(cuts, dtype=float)
    if values.ndim != 2 or cuts.shape != (6, values.shape[1]):
        raise ValueError("分位数阈值与预测目标的维度不一致。")
    if not np.isfinite(values).all() or not np.isfinite(cuts).all() or (np.diff(cuts, axis=0) < 0).any():
        raise ValueError("分级需要有限数值及非递减阈值。")
    return np.sum(values[:, None, :] > cuts[None, :, :], axis=1).astype(np.int8)


def grade_metrics(observed, predicted):
    # 这是同一训练折阈值下的相对等级对比，不是原始人工 0–6 评分的准确率。
    error = np.asarray(predicted, dtype=float) - np.asarray(observed, dtype=float)
    return dict(grade_MAE=float(np.mean(np.abs(error))), grade_RMSE=float(np.sqrt(np.mean(error**2))),
                exact_match_rate=float(np.mean(error == 0)), within_one_rate=float(np.mean(np.abs(error) <= 1)))


def make_summary_frame(meta, targets, grades, columns):
    """按 3 号 summary 列名导出；目标按名称映射，不按 4 号列位置套写。"""
    grades = np.asarray(grades)
    if grades.shape != (len(meta), len(targets)) or not np.isfinite(grades).all():
        raise ValueError("summary 的样本/目标维度或数值不合法。")
    if ((grades < 0) | (grades > 6) | (grades != np.floor(grades))).any():
        raise ValueError("summary 只允许 0–6 整数等级。")
    positions = {name: i for i, name in enumerate(targets)}
    result = {}
    for name in columns:
        target = SUMMARY_ALIASES.get(name, name)
        if name == "Beer":
            result[name] = meta.beer.to_numpy()
        elif name == "comments":
            result[name] = [""] * len(meta)
        elif target in positions:
            result[name] = grades[:, positions[target]].astype(np.int8)
        else:
            # --targets af/flavor 仅试算部分目标，未训练字段留空，不能用 0 冒充预测。
            result[name] = pd.array([pd.NA] * len(meta), dtype="Int64")
    return pd.DataFrame(result, columns=columns)



def require_matplotlib():
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as error:
        raise ImportError("缺少 matplotlib，请安装：python -m pip install matplotlib") from error
    return plt


def resolve_results(results, data_root, output_root=None):
    def compatible(folder):
        if not all((folder / name).is_file() for name in REQUIRED_RESULTS):
            return False
        try:
            info = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8-sig"))
        except (ValueError, OSError):
            return False
        return isinstance(info, dict) and info.get("protocol") == PROTOCOL and info.get("status") == "completed"
    if results not in (None, "latest"):
        folder = Path(results).expanduser().resolve()
        if not compatible(folder):
            raise ValueError(f"不是本版六模型流程的完整结果：{folder}")
        return folder
    parent = Path(output_root or Path(data_root) / "运行结果").resolve()
    folders = [p for p in parent.iterdir() if p.is_dir() and compatible(p)] if parent.is_dir() else []
    if not folders:
        raise FileNotFoundError("尚无本版完整结果。请先运行 python q1_without_joint.py。")
    return max(folders, key=lambda p: ((p / "cv_summary.csv").stat().st_mtime_ns, p.name))


def _save_fig(fig, folder, name, plt):
    fig.savefig(folder / f"{name}.png", dpi=220, bbox_inches="tight")
    fig.savefig(folder / f"{name}.pdf", bbox_inches="tight")
    plt.close(fig)


def plot_results(results, heatmap_beers=None, repeat=1, quantile_fold=1):
    """只画最佳模型评分、六模型指标对比和分位数图，并同步评分工作簿。"""
    results = Path(results).resolve()
    info, summary, best, oof, thresholds = load_result_tables(results)
    beers = list(range(1, 13)) if heatmap_beers is None else list(heatmap_beers)
    frames = build_best_frames(info, best, oof, thresholds, beers, repeat)
    matrix = frames["heatmap"]
    labels = list(matrix.columns)
    targets = [SUMMARY_ALIASES.get(column, column) for column in labels]
    selected_cuts = thresholds[(thresholds.repeat == repeat) & (thresholds.fold == quantile_fold)]
    if selected_cuts.empty or selected_cuts.target.duplicated().any():
        raise ValueError("指定的训练折分位点不存在或存在重复目标。")
    selected_cuts = selected_cuts.set_index("target").loc[targets]
    export_best_outputs(results, info, best, frames, repeat)
    plt = require_matplotlib()
    from matplotlib import font_manager
    from matplotlib.colors import BoundaryNorm
    names = {f.name for f in font_manager.fontManager.ttflist}
    font = next((s for s in ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "SimSun"] if s in names), "DejaVu Sans")
    plt.rcParams.update({"font.family": font, "axes.unicode_minus": False, "font.size": 10,
        "figure.facecolor": "white", "savefig.facecolor": "white", "pdf.fonttype": 42})
    folder = results / "figures"
    folder.mkdir(exist_ok=True)
    cmap = plt.get_cmap("YlGnBu", 7)
    norm = BoundaryNorm(np.arange(-0.5, 7.5), cmap.N)

    # 图 1：只有所选最佳模型，12 行对应已确认的 beer 编号。
    fig, ax = plt.subplots(figsize=(max(15, len(targets)*.42), 6.5), layout="constrained")
    values = matrix.to_numpy(int)
    image = ax.imshow(values, cmap=cmap, norm=norm, aspect="auto", interpolation="nearest")
    ax.set_xticks(range(len(targets)), labels=labels, rotation=60, ha="right", fontsize=8)
    ax.set_yticks(range(12), labels=[f"Beer {b}" for b in beers], fontsize=10)
    for i in range(12):
        for j in range(len(targets)):
            ax.text(j, i, str(values[i, j]), ha="center", va="center", fontsize=8,
                    color="white" if values[i, j] >= 4 else "#222222")
    fig.colorbar(image, ax=ax, ticks=range(7), shrink=.85, pad=.01, label="0–6 相对评分")
    ax.set_title(f"{best['model']}：12 款啤酒的感官评分\n"
        f"按外层 CV 平均 MAE 选取最佳模型；第 {repeat} 次重复的折外评分，与 Excel 一致", fontsize=14, pad=15)
    _save_fig(fig, folder, "01_beer12_score_heatmap", plt)
    heatmap_values = matrix.reset_index().rename(columns={"Beer": "beer"})
    heatmap_values.insert(0, "model", best["model"])
    heatmap_values.to_csv(folder / "beer12_heatmap_values.csv", index=False, encoding="utf-8-sig")

    # 图 2：仍保留六种方法的评价指标，最佳模型用边框标出。
    fig, axes = plt.subplots(1, 3, figsize=(18, 6), layout="constrained")
    table = summary.set_index("model").loc[list(MODELS)]
    colors = ["#467594", "#3D968C", "#CE9855", "#9A73A8", "#C86B67", "#366955"]
    for ax, metric in zip(axes, ("MAE", "RMSE", "R2")):
        values = table[f"{metric}_mean"].to_numpy(float)
        std = table[f"{metric}_std"].fillna(0).to_numpy(float)
        bars = ax.bar(MODELS, values, yerr=std, capsize=4, color=colors, width=.65)
        winner_bar = bars[list(MODELS).index(best["model"])]
        winner_bar.set_edgecolor("#222222")
        winner_bar.set_linewidth(2)
        ax.tick_params(axis="x", labelrotation=30, labelsize=9)
        ax.bar_label(bars, labels=[f"{v:.3f}" if np.isfinite(v) else "N/A" for v in values], padding=5, fontsize=9)
        ax.set_title("$R^2$（越高越好）" if metric == "R2" else f"{metric}（越低越好）")
        ax.axhline(0, color="#777777", linewidth=.7)
        ax.grid(axis="y", alpha=.18)
        ax.margins(y=.25)
        if metric == "R2" and np.isfinite(values-std).any() and np.nanmin(values-std) < -1:
            ax.set_yscale("symlog", linthresh=.25)
            ax.set_ylabel("对称对数刻度，保留负值")
    fig.suptitle("六种方法的外层交叉验证指标\n连续感官标尺、目标等权；误差线为折间标准差", fontsize=14)
    _save_fig(fig, folder, "02_model_metrics_comparison", plt)

    # 图 3：展示同一训练折的分级阈值。
    fig = plt.figure(figsize=(12, max(10, len(targets)*.27+3)), layout="constrained")
    grid = fig.add_gridspec(2, 1, height_ratios=[1.2, max(7, len(targets)*.27)])
    top, ax = fig.add_subplot(grid[0]), fig.add_subplot(grid[1])
    edges = np.arange(8)/7
    for g in range(7):
        top.axvspan(edges[g], edges[g+1], color=cmap(g))
        top.text((edges[g]+edges[g+1])/2, .5, f"评分 {g}", ha="center", va="center",
            color="white" if g >= 4 else "#222222", transform=top.get_xaxis_transform())
    top.set(xlim=(0, 1), yticks=[], xticks=edges)
    top.set_xticklabels(["0"]+[f"{k}/7\n{k/7:.2%}" for k in range(1, 7)]+["1"], fontsize=9)
    top.set_title("0–6 七档对应的分位概率区间（并列值保留，实际人数不必完全相等）", fontsize=11)
    values = selected_cuts[[f"q{k}_of_7" for k in range(1, 7)]].to_numpy(float)
    image = ax.imshow(values, cmap="viridis", aspect="auto", interpolation="nearest")
    ax.set_xticks(range(6), labels=[f"Q{k}/7 ({k/7:.2%})" for k in range(1, 7)], fontsize=9)
    ax.set_yticks(range(len(targets)), labels=labels, fontsize=8)
    lo, hi = float(values.min()), float(values.max())
    for i in range(len(targets)):
        for j in range(6):
            ax.text(j, i, f"{values[i,j]:.3f}", ha="center", va="center", fontsize=8,
                color="white" if values[i,j] < lo+.55*(hi-lo) else "#222222")
    fig.colorbar(image, ax=ax, shrink=.6, label="表格 4 连续感官值")
    fig.suptitle(f"分级阈值：第 {repeat} 次重复、第 {quantile_fold} 个外层训练折（n={int(selected_cuts.n_train.iloc[0])}）\n"
        "只用训练标签计算，边界相等归低档；每一折的全部阈值另存 CSV", fontsize=13)
    _save_fig(fig, folder, "03_quantile_thresholds", plt)
    (folder / "图片说明.txt").write_text(
        f"图1：{best['model']}，beer={beers}，第 {repeat} 次重复的折外整数评分；与评分工作簿同源。\n"
        "图2：六种模型的外层连续值 MAE、宏平均 RMSE、宏平均 R2；折标准差不是置信区间。\n"
        f"图3：第 {repeat} 次重复第 {quantile_fold} 折训练标签的六个分位点。\n"
        "三类图同时保存 PNG/PDF。0–6 为相对分位数等级。\n", encoding="utf-8")
    print(f"三类图片已保存：{folder}", flush=True)
    return folder




def choose_best_model(summary):
    """按全部外层折的 MAE 均值选一个报告模型，不挑选个别测试折。"""
    if summary.model.duplicated().any() or set(summary.model) != set(MODELS):
        raise ValueError("选择最佳模型需要六种方法的完整交叉验证汇总。")
    if not np.isfinite(summary[["MAE_mean", "RMSE_mean"]].to_numpy(float)).all():
        raise ValueError("最佳模型的选择指标存在缺失或非有限值。")
    table = summary.copy()
    table["_r2_order"] = -table.R2_mean.fillna(-np.inf)
    table["_fixed_order"] = table.model.map({name: i for i, name in enumerate(MODELS)})
    winner = table.sort_values(["MAE_mean", "RMSE_mean", "_r2_order", "_fixed_order"]).iloc[0]
    return dict(model=str(winner.model), selection_metric="MAE_mean",
        rule="lowest outer-CV mean MAE; ties: lowest RMSE, highest R2, fixed model order",
        cv_metrics={key: float(winner[key]) if np.isfinite(winner[key]) else None
                    for key in ("MAE_mean", "RMSE_mean", "R2_mean")},
        evaluation_note="reporting choice uses CV comparison; these scores are not an independent test of the final selection")

def require_excel():
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError("缺少 Excel 导出依赖，请安装：python -m pip install openpyxl") from error
    return openpyxl

def load_result_tables(results):
    results = Path(results).resolve()
    info = json.loads((results / "run_manifest.json").read_text(encoding="utf-8-sig"))
    if info.get("protocol") != PROTOCOL or info.get("status") != "completed":
        raise ValueError("只能读取本版完整运行结果。")
    summary = pd.read_csv(results / "cv_summary.csv")
    best = json.loads((results / "best_model.json").read_text(encoding="utf-8"))
    expected = choose_best_model(summary)
    if best.get("model") != expected["model"] or best.get("selection_metric") != expected["selection_metric"]:
        raise ValueError("最佳模型记录与交叉验证汇总不一致。")
    oof = pd.read_csv(results / "oof_predictions.csv", dtype={"beer_id": str})
    thresholds = pd.read_csv(results / "quantile_thresholds.csv")
    if set(oof.model) != set(MODELS):
        raise ValueError("折外预测必须包含全部六种方法。")
    return info, summary, best, oof, thresholds

def build_best_frames(info, best, oof, thresholds, beers, repeat):
    """只为最佳模型重建连续预测和整数评分，热力图与 Excel 共用。"""
    beers = list(beers)
    if len(beers) != 12 or len(set(beers)) != 12:
        raise ValueError("评分热力图必须指定 12 个不同的 beer 编号。")
    targets = info["target_order"]
    sample = oof[(oof.model == best["model"]) & (oof.repeat == repeat)].copy()
    if sample.empty or sample.duplicated(["beer", "target"]).any():
        raise ValueError("最佳模型的指定重复结果为空，或存在重复的样本/目标。")
    if set(sample.target) != set(targets) or not sample.groupby("beer").size().eq(len(targets)).all():
        raise ValueError("最佳模型的感官字段不完整。")
    if sample.beer.nunique() != info["n_samples"]:
        raise ValueError("最佳模型的样本覆盖不完整。")
    if sample.groupby("beer")[["beer_id", "style", "fold"]].nunique(dropna=False).gt(1).any().any():
        raise ValueError("同一样本的编号信息或测试折不一致。")
    qcols = [f"q{k}_of_7" for k in range(1, 7)]
    cuts = thresholds[thresholds.repeat == repeat][["fold", "target"]+qcols]
    if cuts.duplicated(["fold", "target"]).any():
        raise ValueError("同一测试折对应的训练分位点重复。")
    sample = sample.merge(cuts, on=["fold", "target"], how="left", validate="many_to_one")
    values = sample.predicted_continuous.to_numpy(float)
    cut_values = sample[qcols].to_numpy(float)
    if not np.isfinite(values).all() or not np.isfinite(cut_values).all() or (np.diff(cut_values, axis=1) < 0).any():
        raise ValueError("最佳模型预测或训练分位点不合法。")
    sample["predicted_grade_0_6"] = (values[:, None] > cut_values).sum(axis=1).astype(np.int8)
    meta = sample[["beer", "beer_id", "style"]].drop_duplicates("beer").sort_values("beer").reset_index(drop=True)
    meta = meta.rename(columns={"style": "tasting_category_fine"})
    if set(beers)-set(meta.beer):
        raise ValueError("热力图编号在最佳模型的预测中不存在。")
    predicted = sample.pivot(index="beer", columns="target", values="predicted_continuous").loc[meta.beer, targets]
    grades = sample.pivot(index="beer", columns="target", values="predicted_grade_0_6").loc[meta.beer, targets]
    summary_frame = make_summary_frame(meta, targets, grades.to_numpy(), info["summary_columns"])
    beer12 = summary_frame.set_index("Beer").loc[beers].reset_index()
    continuous = pd.concat([meta, predicted.reset_index(drop=True)], axis=1)
    display_columns = [column for column in info["summary_columns"][1:-1]
                       if SUMMARY_ALIASES.get(column, column) in targets]
    heatmap = beer12.set_index("Beer")[display_columns].astype(int)
    return dict(summary=summary_frame, beer12=beer12, continuous=continuous, heatmap=heatmap)

def export_best_outputs(results, info, best, frames, repeat):
    """只导出一个模型、一个明确重复编号的评分，不混合不同重复的分位数等级。"""
    require_excel()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    results = Path(results)
    score_folder, continuous_folder = results / "summary_outputs", results / "continuous_outputs"
    score_folder.mkdir(exist_ok=True)
    continuous_folder.mkdir(exist_ok=True)
    frames["summary"].to_csv(score_folder / "summary_best_model.csv", index=False, encoding="utf-8-sig")
    frames["continuous"].to_csv(continuous_folder / "predicted_table4_best_model.csv", index=False, encoding="utf-8-sig")
    explanation = pd.DataFrame([
        ("模型", best["model"]), ("选取标准", "外层交叉验证 MAE 均值最小"),
        ("交叉验证 MAE", best["cv_metrics"]["MAE_mean"]),
        ("交叉验证 RMSE", best["cv_metrics"]["RMSE_mean"]),
        ("交叉验证 R²", best["cv_metrics"]["R2_mean"]),
        ("评分对应重复编号", repeat), ("summary 样本数", len(frames["summary"])),
        ("已建模感官指标数", len(info["target_order"])),
        ("预测来源", "交叉验证折外预测；样本未参与预测自身的基模型训练"),
        ("评分定义", "各自训练折标签的七等分位数等级，边界相等归低档"),
        ("未训练字段", "留空；不以 0 冒充预测"),
        ("评价范围", "模型通过本次交叉验证比较选出，尚无额外独立测试")], columns=["项目", "内容"])
    destination = score_folder / "best_model_scores.xlsx"
    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        for sheet, frame in (("summary", frames["summary"]), ("beer12", frames["beer12"]),
                             ("continuous", frames["continuous"]), ("说明", explanation)):
            frame.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "B2" if sheet != "说明" else "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 36
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="285F73")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for j, column in enumerate(frame.columns, 1):
                ws.column_dimensions[get_column_letter(j)].width = min(28, max(12, len(str(column))+2))
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell = row[0]
                    if isinstance(cell.value, (float, int)):
                        cell.number_format = "0.0000" if sheet in {"continuous", "说明"} and column != "beer" else "0"
            if sheet == "说明":
                ws.column_dimensions["A"].width = 25
                ws.column_dimensions["B"].width = 80
                for row in ws.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical="top")
                    ws.row_dimensions[row[0].row].height = 30
    print(f"最佳模型 {best['model']} 的评分已保存：{destination}", flush=True)
    return destination

def run_training(args):
    require_excel()
    if not args.no_plots:
        require_matplotlib()
    data = load_data(args)
    n, width = data.Y.shape
    min_outer = n-int(np.ceil(n/args.folds))
    if args.folds > n//2 or args.inner_folds > min_outer//2:
        raise ValueError("折数过多，每个验证折应至少有两个样本。")
    output = Path(args.output_root or args.data_root / "运行结果").resolve() / f"六模型结果_{datetime.now():%Y%m%d_%H%M%S_%f}"
    output.mkdir(parents=True, exist_ok=False)
    manifest = dict(status="running", protocol=PROTOCOL, arguments=vars(args), models=list(MODELS), n_experiments=len(MODELS),
        n_samples=n, n_features=len(data.chemicals), n_targets=width,
        feature_order=data.chemicals, target_order=data.targets, summary_columns=data.summary_columns,
        summary_aliases=SUMMARY_ALIASES, heatmap_beers=args.heatmap_beers,
        tuning=dict(Ridge="per-target alpha, native multioutput Ridge", RF="per-target parameters",
            GBDT="per-target parameters", PLS2="one joint multioutput fit, shared n_components",
            selection_metric=args.selection_metric, inner_repeats=args.inner_repeats,
            aggregation="mean validation-fold score over all inner repetitions",
            shared_splits="identical outer and inner splits and seed schedule for all methods"),
        fusion=dict(base_models=list(BASE_MODELS), selection_metric=args.selection_metric,
            weight_loss="target-standardized MSE + lambda * squared distance to selected-model one-hot",
            nonnegative_weights=True, sum_to_one=True, intercept=False,
            blend_shrinkage=args.blend_shrinkage, shrinkage_selection="fixed before running",
            training_source="selected-candidate inner OOF predictions averaged per sample over inner repeats",
            inner_error_note="reused for tuning and fusion fitting; diagnostic only, not independent validation",
            evaluation="outer held-out predictions only; no selection on outer scores"),
        constant_references="training mean/median only; diagnostics, not candidates or additional experiments",
        grading=dict(probabilities=GRADE_PROBABILITIES.tolist(), method="linear",
            source="outer training labels only; shared by all six methods", equal_boundary="lower grade"),
        preprocessing="training-only median imputation, clipping and scaling of all 231 chemical fields",
        sample_protocol="all table-4 samples enter outer KFold; beer 1-12 shown using their held-out predictions",
        inputs={k: dict(path=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for k, p in data.files.items()},
        versions=dict(python=platform.python_version(), numpy=np.__version__, pandas=pd.__version__, sklearn=sklearn.__version__),
        source_code_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())

    def write_manifest():
        (output / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    write_manifest()
    save_csv([dict(summary_column=c, table4_target=SUMMARY_ALIASES.get(c, c),
        modeled=SUMMARY_ALIASES.get(c, c) in data.targets) for c in data.summary_columns[1:-1]],
        output / "summary_column_mapping.csv")
    oof = {name: np.full((args.repeats, n, width), np.nan) for name in MODELS}
    grades = {name: np.full((args.repeats, n, width), -1, dtype=np.int8) for name in MODELS}
    truth_grades = np.full((args.repeats, n, width), -1, dtype=np.int8)
    membership = np.zeros((args.repeats, n), dtype=int)
    metrics_rows, search_rows, params_rows, split_rows, inner_split_rows = [], [], [], [], []
    clip_rows, threshold_rows, grade_rows, fusion_rows = [], [], [], []
    reference_rows, reference_fold_rows, style_rows = [], [], []
    fusion_oof_folder = output / "fusion_training_oof"
    fusion_oof_folder.mkdir()
    print(f"数据：{n} 个样本，{len(data.chemicals)} 个化学特征，{width} 个感官目标。")
    print(f"六种方法共同参与比较；内层 {args.inner_folds} 折 × {args.inner_repeats} 次。")
    print(f"调参目标：{args.selection_metric}；TargetBlend 用平方损失。结果：{output}", flush=True)
    for repeat in range(1, args.repeats+1):
        outer = KFold(args.folds, shuffle=True, random_state=args.seed+repeat-1)
        for number, (train, valid) in enumerate(outer.split(data.X), 1):
            seed = args.seed+repeat*10000+number*100
            print(f"\n重复 {repeat}，外层折 {number}/{args.folds}，训练 {len(train)} / 测试 {len(valid)}。", flush=True)
            membership[repeat-1, valid] = number
            cuts = fit_grade_cuts(data.Y[train])
            truth_grades[repeat-1, valid] = quantile_grades(data.Y[valid], cuts)
            train_grades = quantile_grades(data.Y[train], cuts)
            for j, target in enumerate(data.targets):
                threshold_rows.append(dict(repeat=repeat, fold=number, target=target, n_train=len(train),
                    unique_cut_count=len(np.unique(cuts[:, j])),
                    **{f"q{k+1}_of_7": float(cuts[k, j]) for k in range(6)},
                    **{f"train_count_grade_{g}": int(np.sum(train_grades[:, j] == g)) for g in range(7)}))
            for role, indices in (("train", train), ("validation", valid)):
                split_rows.extend(dict(repeat=repeat, fold=number, role=role, beer=int(data.meta.iloc[i].beer)) for i in indices)
                counts = data.meta.iloc[indices].tasting_category_fine.value_counts().reindex(
                    sorted(data.meta.tasting_category_fine.unique()), fill_value=0)
                style_rows.extend(dict(repeat=repeat, fold=number, role=role, style=style, n_samples=int(count))
                                  for style, count in counts.items())
            # 划分只生成一次，六种方法原样复用；每次内层重复恰好覆盖训练样本一次。
            inner_splits = []
            inner_membership = np.zeros((args.inner_repeats, len(train)), dtype=int)
            for ir in range(args.inner_repeats):
                inner = KFold(args.inner_folds, shuffle=True, random_state=seed+ir*997)
                current_splits = list(inner.split(train))
                coverage = np.zeros(len(train), dtype=int)
                for i, (it, iv) in enumerate(current_splits, 1):
                    coverage[iv] += 1
                    inner_membership[ir, iv] = i
                    for role, indices in (("train", it), ("validation", iv)):
                        inner_split_rows.extend(dict(repeat=repeat, outer_fold=number, inner_repeat=ir+1,
                            inner_fold=i, role=role, beer=int(data.meta.iloc[train[k]].beer)) for k in indices)
                if not np.all(coverage == 1):
                    raise RuntimeError("每次内层重复必须覆盖所有训练样本一次。")
                inner_splits.extend(current_splits)

            outer_data = prepare_fold(data, train, valid, args)
            inner_data = [prepare_fold(data, train[it], train[iv], args) for it, iv in inner_splits]
            transform = outer_data["transform"]
            clip_rows.extend(dict(repeat=repeat, fold=number, feature=chemical,
                lower=transform.lower[i], upper=transform.upper[i]) for i, chemical in enumerate(data.chemicals))
            inner_predictions, outer_predictions = {}, {}
            for name in BASE_MODELS:
                parameters, parts = select_model(name, inner_data, args, seed, repeat, number,
                                                 data.targets, search_rows, params_rows)
                inner_predictions[name] = assemble_inner_oof(parts, inner_splits, data.Y[train], args.inner_repeats)
                outer_predictions[name] = predict_model(name, parameters, outer_data, args, seed)
            inner_stack = np.stack([inner_predictions[name] for name in BASE_MODELS], axis=-1)
            weights, target_records = learn_target_fusion(data.Y[train], inner_stack, data.targets,
                args.blend_shrinkage, args.selection_metric)
            fusion_rows.extend(dict(repeat=repeat, fold=number, **row) for row in target_records)
            fallback_count = sum(row["optimizer_status"] != "optimized" for row in target_records)
            if fallback_count:
                print(f"  {fallback_count} 个目标的权重优化回退为 TargetSelect，已记录原因。", flush=True)
            outer_stack = np.stack([outer_predictions[name] for name in BASE_MODELS], axis=-1)
            for name in FUSION_MODELS:
                outer_predictions[name] = apply_target_fusion(outer_stack, weights[name])
            fold_labels = [",".join(str(int(v)) for v in inner_membership[:, i]) for i in range(len(train))]
            pd.DataFrame(dict(repeat=repeat, outer_fold=number,
                inner_fold_by_repeat=np.repeat(fold_labels, width), n_inner_predictions=args.inner_repeats,
                beer=np.repeat(data.meta.iloc[train].beer.to_numpy(), width), target=np.tile(data.targets, len(train)),
                observed_continuous=data.Y[train].ravel(),
                **{f"predicted_{name}": inner_predictions[name].ravel() for name in BASE_MODELS})).to_csv(
                    fusion_oof_folder / f"repeat{repeat}_fold{number}.csv", index=False, encoding="utf-8-sig")
            # 冻结所有参数和权重后才计算外层成绩；从不回传成绩修改模型。
            for name in MODELS:
                key = name
                prediction = outer_predictions[name]
                oof[key][repeat-1, valid] = prediction
                prediction_grades = quantile_grades(prediction, cuts)
                grades[key][repeat-1, valid] = prediction_grades
                score, _ = metric_values(data.Y[valid], prediction)
                metrics_rows.append(dict(repeat=repeat, fold=number, model=name, n_train=len(train), n_valid=len(valid),
                    n_model_features=outer_data["x_train"].shape[1],
                    **score))
                grade_rows.append(dict(repeat=repeat, fold=number, model=name,
                    **grade_metrics(truth_grades[repeat-1, valid], prediction_grades)))
                print(f"  {key}: MAE={score['MAE']:.4f}, RMSE={score['RMSE']:.4f}, R2={score['R2']:.4f}", flush=True)

            # 不用化学特征的训练均值/中位数，仅作诊断，不加入六种模型或融合。
            for reference, center in (("Mean", data.Y[train].mean(axis=0)), ("Median", np.median(data.Y[train], axis=0))):
                prediction = np.broadcast_to(center, (len(valid), width))
                score, detail = metric_values(data.Y[valid], prediction)
                reference_fold_rows.append(dict(repeat=repeat, fold=number, reference=reference, **score))
                reference_rows.extend(dict(repeat=repeat, fold=number, reference=reference, target=target,
                    **{metric: float(detail[metric][j]) for metric in ("MAE", "RMSE", "R2")})
                    for j, target in enumerate(data.targets))
            for filename, rows in (("fold_metrics", metrics_rows), ("inner_search", search_rows),
                ("selected_parameters", params_rows), ("cv_splits", split_rows), ("inner_cv_splits", inner_split_rows),
                ("cv_style_counts", style_rows),
                ("clipping_bounds", clip_rows), ("quantile_thresholds", threshold_rows),
                ("quantile_grade_fold_metrics", grade_rows), ("target_fusion_weights", fusion_rows),
                ("constant_reference_per_target", reference_rows), ("constant_reference_fold_metrics", reference_fold_rows)):
                save_csv(rows, output / f"{filename}.csv")
    if (membership == 0).any() or any(not np.isfinite(a).all() for a in oof.values()):
        raise RuntimeError("折外预测未完整覆盖。")
    if (truth_grades < 0).any() or any((a < 0).any() for a in grades.values()):
        raise RuntimeError("整数分级未完整覆盖。")
    frames, overall, target_rows, grade_overall = [], [], [], []
    for name in MODELS:
        key = name
        for r, prediction in enumerate(oof[key]):
            score, detail = metric_values(data.Y, prediction)
            overall.append(dict(model=name, repeat=r+1, **score))
            grade_overall.append(dict(model=name, repeat=r+1, **grade_metrics(truth_grades[r], grades[key][r])))
            for j, target in enumerate(data.targets):
                target_rows.append(dict(model=name, repeat=r+1, target=target,
                    **{metric: float(detail[metric][j]) for metric in ("MAE", "RMSE", "R2")}))
            frames.append(pd.DataFrame(dict(model=name, repeat=r+1, fold=np.repeat(membership[r], width),
                beer=np.repeat(data.meta.beer.to_numpy(), width), beer_id=np.repeat(data.meta.beer_id.to_numpy(), width),
                style=np.repeat(data.meta.tasting_category_fine.to_numpy(), width), target=np.tile(data.targets, n),
                observed_continuous=data.Y.ravel(), predicted_continuous=prediction.ravel())))
    pd.concat(frames, ignore_index=True).to_csv(output / "oof_predictions.csv", index=False, encoding="utf-8-sig")
    for filename, rows in (("oof_metrics_by_repeat", overall), ("per_target_metrics", target_rows),
                           ("quantile_grade_oof_metrics", grade_overall)):
        save_csv(rows, output / f"{filename}.csv")
    keys = ["model"]
    summary = export_summary(pd.DataFrame(metrics_rows), keys, output / "cv_summary.csv")
    best = choose_best_model(summary)
    (output / "best_model.json").write_text(json.dumps(best, ensure_ascii=False, indent=2), encoding="utf-8")
    export_summary(pd.DataFrame(overall), keys, output / "oof_summary.csv")
    export_summary(pd.DataFrame(reference_fold_rows), ["reference"], output / "constant_reference_summary.csv")
    export_fusion_comparison(pd.DataFrame(metrics_rows), output)
    fusion_frame = pd.DataFrame(fusion_rows)
    counts = pd.crosstab(fusion_frame.target, fusion_frame.selected_model).reindex(
        columns=BASE_MODELS, fill_value=0)
    counts.to_csv(output / "target_model_selection_counts.csv", encoding="utf-8-sig")
    weight_summary = fusion_frame.groupby("target", sort=False)[
        [f"weight_{name}" for name in BASE_MODELS]].agg(["mean", "std"])
    weight_summary.columns = [f"{name}_{stat}" for name, stat in weight_summary.columns]
    weight_summary.reset_index().to_csv(output / "target_blend_weight_summary.csv", index=False, encoding="utf-8-sig")
    manifest.update(status="completed", best_model=best["model"], scores_status="pending",
                    plots_status="disabled" if args.no_plots else "pending")
    write_manifest()
    try:
        if args.no_plots:
            tables = load_result_tables(output)
            best_frames = build_best_frames(tables[0], tables[2], tables[3], tables[4], args.heatmap_beers, args.plot_repeat)
            export_best_outputs(output, tables[0], tables[2], best_frames, args.plot_repeat)
        else:
            plot_results(output, args.heatmap_beers, args.plot_repeat, args.quantile_fold)
            manifest["plots_status"] = "completed"
        manifest["scores_status"] = "completed"
    except Exception as error:
        manifest.update(scores_status="failed", report_error=repr(error))
        if not args.no_plots:
            manifest["plots_status"] = "failed"
        print(f"训练结果已保存，评分导出或绘图失败：{error}。修复后用 --plot-only 重试。", flush=True)
    write_manifest()
    print("\n外层 CV 均值与折标准差：\n"+summary.to_string(index=False, float_format=lambda x: f"{x:.5f}"))
    print(f"\n最佳模型：{best['model']}（外层平均 MAE 最小）\n结果目录：{output}")
    return output


def main(argv=None):
    args = parse_args(argv)
    if args.plot_only is not None:
        results = resolve_results(args.plot_only,args.data_root,args.output_root)
        print(f"使用结果目录：{results}",flush=True)
        plot_results(results,args.heatmap_beers,args.plot_repeat,args.quantile_fold)
    else:
        run_training(args)


if __name__ == "__main__":
    main()

